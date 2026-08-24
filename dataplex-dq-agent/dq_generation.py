# Copyright 2026 Google LLC
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#     https://www.apache.org/licenses/LICENSE-2.0
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Dataplex Auto-DQ generation pipeline (streamlit-free).

Profiling retrieval, the policy-tag firewall, description generation held to
the TELUS GenAI prompt standards, Collibra glossary enrichment, the Excel
descriptions store, and DQ-rule YAML rendering. Every side effect is injected:
BigQuery clients come in as parameters, the LLM as a `call_llm(system, user)`
callable, and UI messages go through a `notify(level, text)` callable — so the
whole pipeline is testable without Streamlit and free of module-global state."""
import csv
import json
import logging
import os
import re
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from typing import cast

import google.auth.exceptions
import requests
from google.api_core.exceptions import GoogleAPIError
from google.cloud import bigquery

from dq_core import (
    KNOWN_ERRORS, REFERENCE_DIR, bq_ident, collect_repo_scan_ids, finite_float,
    load_abbreviations, parse_llm_json, pick_audit_column, resolve_scan_id,
    validate_scan_id, yaml_quote,
)

logger = logging.getLogger(__name__)


@dataclass(frozen=True)
class Settings:
    """Sidebar snapshot passed to the pipeline (replaces module globals)."""
    source_project_id: str
    source_dataset_id: str
    bq_location: str
    instance: str            # datahub instance; also the scan-id/job prefix
    governance_dir: str = ""
    fuelix_api_key: str = ""
    model_name: str = ""
    project_id: str = ""     # project/dataset/table holding the scan RESULTS
    dataset_id: str = ""
    profile_table_name: str = ""
    gen_descriptions: bool = False
    use_collibra: bool = False
    collibra_url: str = ""
    collibra_domain: str = ""


# --- Profiling ------------------------------------------------------------------
def get_column_profiles(client, project: str, dataset: str, profile_table: str,
                        source_dataset: str, tables: tuple) -> list:
    """Latest profiling metrics per column for the target tables."""
    project, dataset = bq_ident(project, "project"), bq_ident(dataset, "dataset")
    profile_table = bq_ident(profile_table, "table")
    query = f"""
    WITH RankedProfiles AS (
      SELECT data_source.table_id AS table_name, column_name, column_type,
             column_mode, percent_null, percent_unique, min_value, max_value,
             average_value, standard_deviation, top_n,
             ROW_NUMBER() OVER (PARTITION BY data_source.table_id, column_name
                                ORDER BY job_start_time DESC) AS rank
      FROM `{project}.{dataset}.{profile_table}`
      WHERE data_source.dataset_id = @source_dataset_id
        AND data_source.table_id IN UNNEST(@table_names)
    )
    SELECT * FROM RankedProfiles WHERE rank = 1
    """
    job_config = bigquery.QueryJobConfig(query_parameters=[
        bigquery.ScalarQueryParameter("source_dataset_id", "STRING", source_dataset),
        bigquery.ArrayQueryParameter("table_names", "STRING", list(tables)),
    ])
    return [{**{k: row[k] for k in (
                "table_name", "column_name", "column_type", "column_mode",
                "percent_null", "percent_unique", "min_value", "max_value",
                "average_value", "standard_deviation")},
             "top_n": [{"value": i.get("value"), "count": i.get("count"),
                        "percent": i.get("percent")}
                       for i in (row["top_n"] if isinstance(row["top_n"], list) else [])]}
            for row in client.query(query, job_config=job_config)]


# --- Policy-tag firewall: tagged columns never reach any LLM payload -------------
def fetch_bq_metadata(client, project: str, dataset: str, tables: tuple) -> dict:
    """{table: {"tagged": (dotted policy-tagged paths, ...), "table_description":
    str|None, "columns": {dotted path: description|None}}}. One threaded
    get_table per table (only the Tables API exposes policy tags/descriptions);
    raises on lookup failure so callers fail closed."""
    def walk(fields, prefix=""):
        for f in fields:
            yield (prefix + f.name, f.description or None,
                   bool(f.policy_tags and f.policy_tags.names))
            yield from walk(f.fields or (), prefix + f.name + ".")

    def info(table: str) -> dict:
        t = client.get_table(f"{project}.{dataset}.{table}")
        entries = list(walk(t.schema))
        return {"tagged": tuple(path for path, _d, is_tagged in entries if is_tagged),
                "table_description": t.description or None,
                "columns": {path: d for path, d, _t in entries}}

    with ThreadPoolExecutor(max_workers=8) as pool:
        return dict(zip(tables, pool.map(info, tables)))


def _is_tagged(column: str, tagged_paths: tuple) -> bool:
    """True when the column is a tagged field or nested under/above one."""
    return any(column == t or column.startswith(t + ".") or t.startswith(column + ".")
               for t in tagged_paths)


def _label_list(labels: list) -> str:
    head = ", ".join(f"`{x}`" for x in labels[:60])
    return head + (f" … and {len(labels) - 60} more" if len(labels) > 60 else "")


def get_profiles_for_llm(tables: list, settings: Settings, *, fetch_profiles,
                         fetch_bq_meta, notify) -> tuple[list, dict]:
    """(profiles, tagged): latest profiling rows minus policy-tagged columns —
    the single source of profiling data for LLM prompts — plus the tagged-path
    map for the description firewall and the reports."""
    profiles = fetch_profiles(settings.project_id, settings.dataset_id,
                              settings.profile_table_name,
                              settings.source_dataset_id, tuple(tables),
                              settings.bq_location)
    bq_meta = fetch_bq_meta(settings.source_project_id, settings.source_dataset_id,
                            settings.bq_location, tuple(tables))
    tagged = {t: info["tagged"] for t, info in bq_meta.items()}
    kept = [p for p in profiles
            if not _is_tagged(str(p.get("column_name")),
                              tagged.get(str(p.get("table_name")), ()))]
    excluded = [f"{t}.{c}" for t in sorted(tagged) for c in tagged[t]]
    if excluded:
        notify("info", "Policy-tagged columns excluded from every LLM payload (no sample "
               "data, no profiling statistics, no descriptions): "
               + _label_list(excluded))
    return kept, tagged


def existing_descriptions_from(bq_meta: dict) -> dict:
    """bq_meta reshaped to the descriptions shape used everywhere; policy-tagged
    columns are nulled so the firewall holds even when the current metadata is
    chosen for the next steps."""
    return {table: {
        "table_description": info["table_description"],
        "columns": {c: (None if _is_tagged(c, info["tagged"]) else d)
                    for c, d in info["columns"].items()}}
        for table, info in bq_meta.items()}


# --- Optional description generation (TELUS GenAI prompt standards) --------------
SAMPLE_ROW_COUNT = 1000     # window of most recent rows fetched per table
_EVIDENCE_ROWS = 20         # full rows from that window forwarded to the model
_EVIDENCE_TOP_VALUES = 12   # most frequent values per column forwarded
_EVIDENCE_CELL_CAP = 80     # cell truncation before prompting
UNABLE_PHRASE = "Unable to Generate Description"  # LLM sentinel; stored as null
_UNABLE_RE = re.compile(re.escape(UNABLE_PHRASE) + r"\.?", re.IGNORECASE)


def _null_if_unable(text) -> str | None:
    """Stored value for one LLM description: None for blank or the sentinel."""
    text = str(text or "").strip()
    return None if not text or _UNABLE_RE.fullmatch(text) else text


def _latest_partition_filter(client, project: str, dataset: str, table: str,
                             meta: dict) -> tuple[str, list]:
    """(WHERE clause, typed query parameters) pinning the newest partition via
    a zero-scan INFORMATION_SCHEMA.PARTITIONS lookup; ('', []) unless day
    granularity resolves. The day bounds travel as query parameters typed to
    the partition column, so no value is ever spliced into the SQL text."""
    pc, pct = meta["partition_column"], meta["partition_column_type"]
    if not pc:
        return "", []
    project, dataset = bq_ident(project, "project"), bq_ident(dataset, "dataset")
    query = f"""
    SELECT partition_id
    FROM `{project}.{dataset}.INFORMATION_SCHEMA.PARTITIONS`
    WHERE table_name = @table AND partition_id IS NOT NULL
      AND partition_id NOT IN ('__NULL__', '__UNPARTITIONED__')
    ORDER BY partition_id DESC LIMIT 1
    """
    try:
        rows = list(client.query(
            query, job_config=bigquery.QueryJobConfig(query_parameters=[
                bigquery.ScalarQueryParameter("table", "STRING", table)])))
    except GoogleAPIError as e:
        logger.info("partition lookup failed for %s.%s.%s — sampling without "
                    "a partition filter: %s", project, dataset, table, e)
        return "", []
    pid = str(rows[0]["partition_id"]) if rows else ""
    if len(pid) != 8 or not pid.isdigit():
        return "", []
    try:
        day = date.fromisoformat(f"{pid[:4]}-{pid[4:6]}-{pid[6:]}")
    except ValueError:
        return "", []
    pc = bq_ident(pc, "column")
    if pct == "DATE":
        return (f"WHERE `{pc}` = @part_day",
                [bigquery.ScalarQueryParameter("part_day", "DATE", day)])
    if pct in ("TIMESTAMP", "DATETIME"):
        # Typed to the column so the day bounds compare exactly like the
        # string literals BigQuery used to coerce.
        start = datetime(day.year, day.month, day.day,
                         tzinfo=timezone.utc if pct == "TIMESTAMP" else None)
        return (f"WHERE `{pc}` >= @part_start AND `{pc}` < @part_end",
                [bigquery.ScalarQueryParameter("part_start", pct, start),
                 bigquery.ScalarQueryParameter("part_end", pct, start + timedelta(days=1))])
    return "", []


def fetch_last_rows(client, project: str, dataset: str, table: str, meta: dict,
                    n: int = SAMPLE_ROW_COUNT, exclude: frozenset = frozenset(),
                    *, notify) -> list:
    """Most recent n rows as dicts (best effort): newest partition, ordered by
    the audit/temporal column; retried without ORDER BY, then without the
    partition filter (unless the table requires one); [] when every attempt
    fails. `exclude` (policy-tagged top-level columns) is dropped in the SELECT
    itself, so protected data is never even fetched."""
    where, where_params = _latest_partition_filter(client, project, dataset, table, meta)
    order_col = pick_audit_column(meta["temporal_columns"])
    if not order_col and meta["partition_column_type"] in ("DATE", "TIMESTAMP", "DATETIME"):
        order_col = meta["partition_column"]
    order = f" ORDER BY `{bq_ident(order_col, 'column')}` DESC" if order_col else ""
    select = ("* EXCEPT (" + ", ".join(f"`{bq_ident(c, 'column')}`" for c in sorted(exclude))
              + ")" if exclude else "*")
    base = (f"SELECT {select} FROM `{bq_ident(project, 'project')}."
            f"{bq_ident(dataset, 'dataset')}.{bq_ident(table, 'table')}`")
    attempts = {  # most-specific first, deduped when where/order are ''
        f"{base}{' ' + w if w else ''}{o} LIMIT @n": (list(where_params) if w else [])
        for w, o in ((where, order), (where, ""), ("", order), ("", ""))
        if w or not meta["require_partition_filter"]}
    for sql, params in attempts.items():
        job_config = bigquery.QueryJobConfig(query_parameters=[
            *params, bigquery.ScalarQueryParameter("n", "INT64", n)])
        try:
            return [dict(row) for row in client.query(sql, job_config=job_config)]
        except GoogleAPIError as e:
            logger.info("sample attempt for %s failed (%s) — trying a simpler "
                        "query", table, type(e).__name__)
            continue
    notify("warning", f"Could not sample rows from `{table}` — its descriptions will rely "
           "on profiling statistics and names only.")
    return []


def build_sample_evidence(rows: list, max_rows: int = _EVIDENCE_ROWS,
                          top_values: int = _EVIDENCE_TOP_VALUES,
                          cell_cap: int = _EVIDENCE_CELL_CAP) -> dict:
    """Compact prompt-safe digest: per-column non-null counts and most frequent
    values (over all rows), plus the newest max_rows full rows, cells truncated."""
    if not rows:
        return {}

    def clip(value):
        text = "" if value is None else str(value)
        return text[:cell_cap] + ("..." if len(text) > cell_cap else "")

    columns = {}
    for name in rows[0]:
        values = [row.get(name) for row in rows]
        top = Counter(clip(v) for v in values if v is not None).most_common(top_values)
        columns[name] = {"non_null": sum(v is not None for v in values),
                         "rows_sampled": len(values),
                         "top_values": [{"value": k, "count": c} for k, c in top]}
    return {"columns": columns,
            "recent_rows": [{k: clip(v) for k, v in row.items()} for row in rows[:max_rows]]}


# --- Optional Collibra glossary enrichment ----------------------------------------
COLLIBRA_SECRET_PROJECT = "cto-collibra-insights-pr-2267"
COLLIBRA_SECRET_NAME = "collibra_api_key"
COLLIBRA_CSV = os.path.join(REFERENCE_DIR, "collibra_glossary.csv")
_COLLIBRA_CSV_FIELDS = (
    "Full Name", "Name", "Asset Id", "Asset Type", 
    "[Business Term] has acronym [Acronym] > Name",
    "[Business Term] has acronym [Acronym] > Full Name",
    "[Business Term] has acronym [Acronym] > Asset Type",
    "[Business Term] has acronym [Acronym] > Community",
    "[Business Term] has acronym [Acronym] > Domain Type",
    "[Business Term] has acronym [Acronym] > Domain",
    "[Business Term] has acronym [Acronym] > Domain Id",
    "[Business Term] has acronym [Acronym] > Asset Id",
    "Definition", "Status", "Domain", "Community", "Domain Type", "Domain Id"
)
_GLOSSARY_MAX_LINES = 40
_CONTEXT_CHAR_CAP = 80_000  # keep description prompts inside gateway limits
_COLLIBRA_ERRORS = (GoogleAPIError, google.auth.exceptions.GoogleAuthError,
                    requests.RequestException, RuntimeError, ValueError,
                    KeyError, ImportError)


def get_collibra_auth_key(project: str = COLLIBRA_SECRET_PROJECT,
                          secret: str = COLLIBRA_SECRET_NAME) -> str:
    """auth_key from Google Cloud Secret Manager (latest version). Fetched per
    run and held only in locals — credentials are never cached."""
    from google.cloud import secretmanager
    client = secretmanager.SecretManagerServiceClient()
    name = f"projects/{project}/secrets/{secret}/versions/latest"
    return client.access_secret_version({"name": name}).payload.data.decode("UTF-8")


def _collibra_auth_options(base_url: str, key: str, http=requests) -> list:
    """Candidate requests-kwargs for the unknown key scheme, in probe order:
    Bearer, HTTP Basic (user:pass), pre-encoded Basic, session-login cookies."""
    options: list[dict] = [{"headers": {"Authorization": f"Bearer {key}"}},
                           {"headers": {"Authorization": f"Basic {key}"}}]
    if ":" in key:
        user, password = key.split(":", 1)
        options.insert(1, {"auth": (user, password)})
        try:
            resp = http.post(f"{base_url.rstrip('/')}/rest/2.0/auth/sessions",
                             json={"username": user, "password": password},
                             timeout=30)
            if resp.ok and resp.cookies:
                options.append({"cookies": resp.cookies.get_dict()})
        except requests.RequestException as e:
            logger.debug("collibra session probe failed: %s", e)
    return options


def _pick_collibra_auth(base_url: str, auth_key: str, http=requests) -> dict:
    """First auth scheme the API accepts, probed with a 1-row read."""
    statuses = []
    for kwargs in _collibra_auth_options(base_url, auth_key, http):
        resp = http.get(f"{base_url.rstrip('/')}/rest/2.0/assets",
                        params={"limit": 1}, timeout=30, **kwargs)
        if resp.ok:
            return kwargs
        statuses.append(str(resp.status_code))
    raise RuntimeError("Collibra rejected every auth scheme "
                       f"(Bearer/Basic/session): HTTP {', '.join(statuses)}")


def _read_collibra_csv(path: str) -> list[dict]:
    """Rows of a saved glossary CSV; raises on read/parse problems."""
    with open(path, "r", encoding="utf-8", newline="") as fh:
        return [{f: str(row.get(f) or "") for f in _COLLIBRA_CSV_FIELDS}
                for row in csv.DictReader(fh)]


def save_collibra_csv(path: str, glossary: dict, defs: dict) -> None:
    """Replace the saved glossary CSV with `glossary` (one row per term, sorted).
    The definition column takes `defs[asset_id]` and falls back to the previous
    file's value, so definitions accumulated over runs survive a refresh."""
    try:
        carried = {}
        if os.path.exists(path):
            carried = {r["Asset Id"]: r["Definition"]
                       for r in _read_collibra_csv(path) if r.get("Definition")}
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        with open(path, "w", encoding="utf-8", newline="") as fh:
            writer = csv.writer(fh)
            writer.writerow(_COLLIBRA_CSV_FIELDS)
            for term in sorted(glossary):
                m = glossary[term]
                row_data = {f: "" for f in _COLLIBRA_CSV_FIELDS}
                row_data["Name"] = m["name"]
                row_data["Full Name"] = m["full"]
                row_data["Asset Id"] = m["id"]
                row_data["Asset Type"] = "Business Term"
                row_data["Definition"] = defs.get(m["id"]) or carried.get(m["id"], "")
                
                writer.writerow([row_data[col] for col in _COLLIBRA_CSV_FIELDS])
        logger.info("collibra glossary saved: %s (%d terms)", path, len(glossary))
    except (OSError, csv.Error, KeyError, TypeError, UnicodeDecodeError) as e:
        logger.warning("collibra glossary CSV not saved (%s): %s", path, e)


def load_collibra_csv(path: str):
    """(glossary, fetch_defs) from the saved CSV — the same contract as
    load_collibra, with definitions limited to those already saved and zero
    Collibra traffic. Raises RuntimeError on an unreadable file."""
    try:
        rows = _read_collibra_csv(path)
    except (OSError, csv.Error, UnicodeDecodeError) as e:
        raise RuntimeError(f"saved Collibra glossary unreadable ({path}): {e}") from e
        
    glossary = {r["Name"].lower(): {"id": r["Asset Id"], "name": r["Name"],
                                    "full": r["Full Name"] or r["Name"]}
                for r in rows if r.get("Name")}
    defs = {r["Asset Id"]: r["Definition"]
            for r in rows if r.get("Asset Id") and r.get("Definition")}

    def fetch_defs(asset_ids: tuple) -> dict:
        return {a: defs[a] for a in asset_ids if defs.get(a)}
        
    logger.info("collibra glossary loaded from CSV: %s (%d terms)", path, len(glossary))
    return glossary, fetch_defs


def load_collibra(settings: Settings, *, get_key=get_collibra_auth_key, 
                  http=requests, save_path: str = ""):
    """Fetches the Collibra glossary directly from the Collibra REST API v2.
    Replicates the BigQuery join by fetching relations, assets, and attributes."""
    
    auth = _pick_collibra_auth(settings.collibra_url, get_key(), http)
    base_url = settings.collibra_url.rstrip('/')
    
    # 1. Find the "is acronym for" Relation Type
    resp = http.get(f"{base_url}/rest/2.0/relationTypes", timeout=30, **auth)
    resp.raise_for_status()
    relation_type_id = None
    is_acronym_role = True 
    
    for rtype in resp.json().get("results", []):
        if str(rtype.get("role", "")).lower() == "is acronym for":
            relation_type_id = rtype["id"]
            is_acronym_role = True
            break
        elif str(rtype.get("coRole", "")).lower() == "is acronym for":
            relation_type_id = rtype["id"]
            is_acronym_role = False
            break
            
    if not relation_type_id:
        raise RuntimeError("Could not find relation type 'is acronym for' in Collibra.")

    # 2. Fetch all relations of this type
    relations = []
    offset = 0
    while True:
        resp = http.get(f"{base_url}/rest/2.0/relations", 
                        params={"relationTypeId": relation_type_id, "limit": 1000, "offset": offset}, 
                        timeout=60, **auth)
        resp.raise_for_status()
        page = resp.json().get("results", [])
        if not page:
            break
        relations.extend(page)
        offset += len(page)
        
    if not relations:
        return {}, lambda ids: {}
        
    # Extract unique asset IDs to query
    asset_ids = set()
    term_ids = set()
    for rel in relations:
        asset_ids.add(rel["source"]["id"])
        asset_ids.add(rel["target"]["id"])
        # Track which side is the actual business term so we only fetch definitions for those
        term_id = rel["target"]["id"] if is_acronym_role else rel["source"]["id"]
        term_ids.add(term_id)
        
    # 3. Concurrently fetch Asset Names
    asset_names = {}
    def fetch_asset(aid):
        r = http.get(f"{base_url}/rest/2.0/assets/{aid}", timeout=30, **auth)
        if r.ok:
            return aid, r.json().get("displayName") or r.json().get("name")
        return aid, None
        
    with ThreadPoolExecutor(max_workers=10) as pool:
        for aid, name in pool.map(fetch_asset, list(asset_ids)):
            if name:
                asset_names[aid] = name
                
    # 4. Concurrently fetch Definitions for the Business Terms
    defs = {}
    def fetch_def(tid):
        r = http.get(f"{base_url}/rest/2.0/attributes", 
                     params={"assetId": tid, "limit": 20}, timeout=30, **auth)
        if r.ok:
            for attr in r.json().get("results", []):
                tname = str((attr.get("type") or {}).get("name") or "").lower()
                if "definition" in tname or "description" in tname:
                    text = re.sub(r"<[^>]+>", " ", str(attr.get("value") or ""))
                    return tid, re.sub(r"\s+", " ", text).strip()
        return tid, ""
        
    with ThreadPoolExecutor(max_workers=10) as pool:
        for tid, d in pool.map(fetch_def, list(term_ids)):
            if d:
                defs[tid] = d

    # 5. Build the expected Glossary structure
    glossary = {}
    for rel in relations:
        if is_acronym_role:
            acronym_id = rel["source"]["id"]
            term_id = rel["target"]["id"]
        else:
            acronym_id = rel["target"]["id"]
            term_id = rel["source"]["id"]
            
        acronym = asset_names.get(acronym_id, "")
        business_name = asset_names.get(term_id, acronym)
        
        if acronym and term_id:
            term_key = acronym.lower()
            glossary[term_key] = {"id": term_id, "name": acronym, "full": business_name}
            
    if save_path:
        save_collibra_csv(save_path, glossary, defs)
        
    def fetch_defs(asset_ids: tuple) -> dict:
        return {a: defs[a] for a in asset_ids if defs.get(a)}
        
    return glossary, fetch_defs


def _name_tokens(names) -> set:
    """Alphanumeric segments (len >= 2) of the given table/column names."""
    return {t for n in names for t in re.split(r"[^a-zA-Z0-9]+", n.lower()) if len(t) >= 2}


def collibra_hint(glossary: dict, fetch_defs, *names: str) -> str:
    """Up to _GLOSSARY_MAX_LINES 'TERM = Full Name — Definition' lines for
    glossary terms matching the given name segments; '' if none."""
    matched = [glossary[t] for t in sorted(_name_tokens(names))
               if t in glossary][:_GLOSSARY_MAX_LINES]
    if not matched:
        return ""
    defs = fetch_defs(tuple(m["id"] for m in matched if m["id"]))
    return "\n".join(f"{m['name']} = {m['full']}"
                     + (f" — {defs[m['id']]}" if defs.get(m["id"]) else "")
                     for m in matched)


def _reverse_abbreviations(abbrev: dict | None = None) -> dict:
    """{abbreviation: 'full word(s)'} from abbreviations.csv; drop-tokens
    (empty abbreviation) skipped, shared abbreviations joined with ' / '."""
    reverse: dict = {}
    for full, abbr in (load_abbreviations() if abbrev is None else abbrev).items():
        if abbr:
            reverse[abbr] = f"{reverse[abbr]} / {full}" if abbr in reverse else full
    return reverse


def abbrev_hint(*names: str, abbrev: dict | None = None) -> str:
    """Up to _GLOSSARY_MAX_LINES 'abbr = full word(s)' lines for abbreviations
    matching the given name segments; '' if none."""
    reverse = _reverse_abbreviations(abbrev)
    matched = [f"{t} = {reverse[t]}" for t in sorted(_name_tokens(names)) if t in reverse]
    return "\n".join(matched[:_GLOSSARY_MAX_LINES])


# Constraints below follow the TELUS GenAI description prompt standards
# (table-level: 16 constraints; column-level: 15 constraints, batch-adapted).
_SYS_TABLE_DESC = """You are a TELUS data steward generating a business description for a BigQuery table.
You will receive the table name, Dataplex column profiling statistics, and evidence sampled from the table's most recent rows. Use them only to infer business meaning.
If you are unable to generate a description, respond with exactly: Unable to Generate Description

Here are examples of good table descriptions to use as a reference:
Example 1: This table captures the different types of equipment classification. A product equipment classification is a secondary level grouping of the equipment, organized by its use.
Example 2: This table holds outgoing notifications to the TELUS customers for different marketing campaigns.

Your response MUST satisfy ALL the following constraints:
1. The response MUST NOT repeat the provided table name within the description.
2. The response MUST start with the following words: This table contains.
3. The response MUST be clear, concise, and complete.
4. The response MUST explain any key business terms as part of the description.
5. The response MUST be stated in the present tense, in the form of a complete descriptive sentence.
6. The response MUST NOT include jargon or abbreviations.
7. If the response references an acronym, the acronym MUST be spelled out.
8. The response MUST be provided from a business perspective, independent from any database or software context.
9. The response MUST NOT include any potential example/sample values.
10. The response MUST NOT include any details about potential attributes or columns collected or stored in the table.
11. The response MUST EXCLUDE any pleasantries.
12. The response MUST EXCLUDE any external references or citations.
13. The response MUST be confident. DO NOT use words such as "this table may contain" or "this table is likely to contain".
14. The response MUST NOT provide an opinion as to why this information is valuable.
15. The response MUST NOT include filler words and phrases, such as "in the organization" or any synonyms. If this information is required for a complete definition, use TELUS rather than a generic form such as organization or company.
16. If the response references the company name TELUS, ensure that TELUS is capitalized.

When generating the description, also consider:
- Some tables start with the prefix bq_; ignore the prefix and use the rest of the name (bq_customer_info -> customer_info).
- Some tables end with the suffix _dim; ignore the suffix (team_member_dim -> team_member).
- Table names are never about an individual person, so avoid references to specific individuals.
- The profiling statistics and sampled values are context only and MUST NOT appear in the output.
Return ONLY the description text (or the exact fallback phrase) — no preamble, no markdown."""

_SYS_COL_DESC = """You are a TELUS data steward generating business descriptions for the columns of a BigQuery table.
You will receive the table name, its generated table description, the list of columns to describe, Dataplex column profiling statistics, and evidence sampled from the table's most recent rows. Use them only to infer business meaning.
Return ONLY one JSON object (no markdown, no prose) of this shape, covering EVERY listed column:
{"columns": {"<column_name>": "<description>"}}
Use the exact phrase "Unable to Generate Description" as the value for any column you cannot describe.

Here are examples of good column descriptions to use as a reference:
Example 1: This column contains the cardinal direction from which the wind blows in an abbreviated form. Wind directions are always expressed as from whence the wind blows meaning that a North wind blows from North to South.
Example 2: The date printed on the Bill Document indicating the extraction date of Customer data from the billing system.
Example 3: A yes/no indicator identifying whether or not the party is hearing impaired. This field only applies to Parties that are Individuals.
Example 4: The extended customer name such as trade name for the sole proprietorship business customer. It is the trade/firm name that followed the operating type such as Operating As (O/A), Doing Business as (D/B).
Example 5: A code used to represent the discount level that the customer is entitled to for various products and services. This only applies to those customers that are also team members.
Example 6: This column contains an ISO defined unique identifier of a currency. In this instance, it represents the currency in which the customers' invoices are presented.

Every description MUST satisfy ALL the following constraints:
1. The description MUST NOT repeat the provided table name or column name.
2. The description MUST start with the following words: This column contains.
3. The description MUST be clear, concise, and complete.
4. The description MUST explain any key business terms.
5. The description MUST be stated in the present tense, in the form of a complete descriptive sentence.
6. The description MUST NOT include jargon or abbreviations.
7. If the description references an acronym, the acronym MUST be spelled out.
8. The description MUST be provided from a business perspective, independent from any database or software context.
9. The description MUST NOT include any potential example/sample values.
10. The description MUST EXCLUDE any pleasantries.
11. The description MUST EXCLUDE any external references or citations.
12. The description MUST be confident. DO NOT use words such as "this column may contain" or "this column is likely to contain".
13. The description MUST NOT provide an opinion as to why this information is valuable.
14. The description MUST NOT include filler words and phrases, such as "in the organization" or any synonyms. If this information is required for a complete definition, use TELUS rather than a generic form such as organization or company.
15. If the description references the company name TELUS, ensure that TELUS is capitalized.

When generating the column descriptions, also consider:
- Columns ending with the suffix _ind are usually indicators (e.g. INDVDL_HEARING_IMPAIRED_IND indicates whether or not the individual is hearing impaired).
- The profiling statistics and sampled values are context only and MUST NOT appear in the output."""


def _table_profiles(profiles: list, table: str) -> list:
    return [{k: v for k, v in p.items() if k != "table_name"}
            for p in profiles if p.get("table_name") == table]


def _jsonc(obj) -> str:  # compact prompt JSON; default=str for dates/decimals
    return json.dumps(obj, separators=(",", ":"), default=str)


def generate_descriptions(tables: list, profiles: list, tagged: dict,
                          settings: Settings, *, call_llm, fetch_table_meta,
                          fetch_rows, notify, load_glossary=None,
                          abbrev: dict | None = None,
                          existing: dict | None = None,
                          fill_only_missing: bool = False) -> dict:
    """{table: {"table_description": str|None, "columns": {col: str|None}}}:
    two FuelIX calls per table grounded in the last SAMPLE_ROW_COUNT rows and
    the policy-filtered profiling stats. `existing` (already policy-nulled)
    rides along as reference context; fill_only_missing keeps it verbatim and
    generates only the gaps (fully described tables skip sampling and the LLM).
    Tagged columns, blank/sentinel replies and per-table failures store null.
    All side effects are injected: `call_llm(system, user)`, `fetch_table_meta`,
    `fetch_rows(table, meta, exclude)`, `load_glossary()` and `notify`."""
    existing = existing or {}
    meta = fetch_table_meta(settings.source_project_id, settings.source_dataset_id,
                            settings.bq_location, tuple(tables))
    glossary, fetch_defs = {}, None
    if settings.use_collibra and load_glossary is not None:
        try:
            glossary, fetch_defs = load_glossary()
            notify("caption", f"Collibra glossary imported: {len(glossary)} business terms.")
        except _COLLIBRA_ERRORS as e:
            logger.warning("collibra glossary load failed: %s", e)
            notify("warning", f"Collibra glossary unavailable ({e}) — continuing without it.")
    out = {}
    for table in tables:
        tagged_paths = tagged.get(table, ())
        cur = existing.get(table) or {"table_description": None, "columns": {}}
        cur_cols = {c: d for c, d in cur["columns"].items() if d}
        if fill_only_missing and cur["table_description"] and not any(
                not d and not _is_tagged(c, tagged_paths)
                for c, d in cur["columns"].items()):
            out[table] = {"table_description": cur["table_description"],
                          "columns": dict(cur["columns"])}  # fully described: no LLM
            continue
        rows = fetch_rows(table, meta[table],
                          frozenset(t.split(".")[0] for t in tagged_paths))
        stats = _table_profiles(profiles, table)
        stats_json = _jsonc(stats)
        evidence = build_sample_evidence(rows)
        col_names = (list(evidence["columns"]) if evidence
                     else sorted({str(s["column_name"]) for s in stats}))

        def render_context(ev, sj):
            return (f"TABLE NAME: {table}\nPROFILING STATISTICS:\n{sj}\n"
                    f"SAMPLE EVIDENCE (from the last {SAMPLE_ROW_COUNT} rows):\n"
                    + (_jsonc(ev) if ev else "(no rows sampled)"))

        # Glossary/reference blocks are built first and counted against the cap
        extras = ""
        hint = abbrev_hint(table, *col_names, abbrev=abbrev)
        if hint:
            extras += ("\nTELUS ABBREVIATION GLOSSARY (abbreviations.csv; matched "
                       "to this table's name/column segments):\n" + hint)
        ch = ""
        if glossary and fetch_defs is not None:
            try:
                ch = collibra_hint(glossary, fetch_defs, table, *col_names)
            except _COLLIBRA_ERRORS as e:  # degrade once; the run continues
                logger.warning("collibra became unavailable mid-run: %s", e)
                notify("warning", f"Collibra glossary unavailable ({e}) — continuing without it.")
                glossary, fetch_defs = {}, None
        if ch:
            extras += ("\nTELUS BUSINESS GLOSSARY (Collibra; matched to this "
                       "table's name/column segments — reference hints for "
                       "interpreting abbreviations, not data):\n" + ch)
        ref = {k: v for k, v in (("table_description", cur["table_description"]),
                                 ("columns", cur_cols)) if v}
        if ref:
            extras += ("\nEXISTING BIGQUERY DESCRIPTIONS (current metadata — "
                       "reference only, may be incomplete or outdated):\n"
                       + _jsonc(ref)[:20_000])
        budget = max(_CONTEXT_CHAR_CAP - len(extras), _CONTEXT_CHAR_CAP // 4)
        context = render_context(evidence, stats_json)
        if len(context) > budget:
            evidence = build_sample_evidence(rows, max_rows=5, top_values=5, cell_cap=40)
            context = render_context(evidence, stats_json)
        if len(context) > budget and evidence:
            evidence.pop("recent_rows", None)
            context = render_context(evidence, stats_json)
        if len(context) > budget:  # stats only, then without top_n detail
            context = render_context({}, stats_json)
            if len(context) > budget:
                slim = _jsonc([{k: v for k, v in s.items() if k != "top_n"}
                               for s in stats])
                context = render_context({}, slim)
        if len(context) > budget:
            context = context[:budget]  # last resort
            notify("warning", f"Profiling statistics for `{table}` exceed the "
                   "prompt budget even without sample evidence — truncated for "
                   "the LLM call.")
        context += extras
        cols_to_ask = ([c for c in col_names if not cur_cols.get(c)]
                       if fill_only_missing else col_names)
        try:
            table_desc = (cur["table_description"]
                          if fill_only_missing and cur["table_description"]
                          else _null_if_unable(call_llm(
                              _SYS_TABLE_DESC,
                              "Generate the table description.\n" + context)))
            col_map: dict[str, str | None] = {}
            if cols_to_ask:
                col_prompt = (f"COLUMNS TO DESCRIBE: {json.dumps(cols_to_ask)}\n"
                              f"TABLE DESCRIPTION: {table_desc or ''}\n" + context +
                              "\nGenerate the JSON object of column descriptions.")
                data = parse_llm_json(call_llm(_SYS_COL_DESC, col_prompt))
                found = data.get("columns", data) if isinstance(data, dict) else {}
                col_map = {c: _null_if_unable(found.get(c)) for c in cols_to_ask}
        except KNOWN_ERRORS as e:
            logger.exception("description generation failed for table %s", table)
            notify("warning", f"Descriptions for `{table}` failed: {e}")
            table_desc = cur["table_description"] if fill_only_missing else None
            col_map = dict.fromkeys(cols_to_ask)
        if fill_only_missing:  # keep the current descriptions verbatim
            for c, d in cur_cols.items():
                col_map.setdefault(c, d)
        for path in tagged_paths:  # visible in Excel/report, never sent to the LLM
            col_map.setdefault(path, None)
        out[table] = {"table_description": table_desc, "columns": col_map}
    return out


_XLSX_HEADER = ["dataset", "table_name", "column_name", "description", "model",
                "generated_ts"]


def descriptions_xlsx_path(settings: Settings) -> str:
    return os.path.join(os.path.dirname(__file__), "descriptions",
                        f"{settings.instance}_dq_descriptions_{settings.source_dataset_id}.xlsx")


def save_descriptions_xlsx(path: str, descriptions: dict, *, dataset: str,
                           model: str, notify) -> bool:
    """Upsert into the repo Excel keyed on (dataset, table_name, column_name):
    one table-level row (empty column_name) plus one row per column. Rows for
    other datasets/tables are preserved and duplicate key rows all kept in sync.
    Targets the 'descriptions' sheet by name, validates its header, writes
    atomically; returns False (with an error notification) instead of raising
    on I/O problems such as the workbook being open in Excel."""
    try:
        import openpyxl
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        notify("error", "openpyxl is not installed — run `pip install openpyxl` to save "
               "descriptions to Excel.")
        return False

    def clean(value):
        # None passes through so a null description stays a truly blank cell
        # (and the upsert clears any previously saved placeholder text).
        return None if value is None else ILLEGAL_CHARACTERS_RE.sub("", str(value))

    def cell_key(value) -> str:
        return "" if value is None else str(value).strip()

    ts = datetime.now(timezone.utc).isoformat(timespec="seconds")
    fresh = {(cell_key(dataset), cell_key(table), cell_key(col)): clean(desc)
             for table, info in descriptions.items()
             for col, desc in [("", info["table_description"]), *info["columns"].items()]}
    tmp = path + ".tmp"
    try:
        exists = os.path.exists(path)
        wb = openpyxl.load_workbook(path) if exists else openpyxl.Workbook()
        if not exists and wb.active is not None:
            wb.remove(wb.active)
        if "descriptions" in wb.sheetnames:
            ws = wb["descriptions"]
            header = [c.value for c in next(ws.iter_rows(
                min_row=1, max_row=1, max_col=len(_XLSX_HEADER)))]
            if all(v is None for v in header):
                for i, name in enumerate(_XLSX_HEADER, start=1):
                    ws.cell(row=1, column=i, value=name)
            elif header != _XLSX_HEADER:
                notify("error", f"Sheet 'descriptions' in `{path}` has an unexpected "
                       f"header {header} — expected {_XLSX_HEADER}. Fix or "
                       "delete the file, then regenerate.")
                return False
        else:
            ws = wb.create_sheet("descriptions")
            ws.append(_XLSX_HEADER)
        updated = set()
        for row in ws.iter_rows(min_row=2, max_col=len(_XLSX_HEADER)):
            key = (cell_key(row[0].value), cell_key(row[1].value), cell_key(row[2].value))
            if key in fresh:
                # assign via .value: ws.cell(value=None) skips the write and
                # would leave stale text behind instead of blanking the cell.
                cells = cast("tuple[Cell, ...]", row)
                cells[3].value, cells[4].value, cells[5].value = fresh[key], model, ts
                updated.add(key)
        for (ds, table, col), desc in fresh.items():
            if (ds, table, col) not in updated:
                ws.append([ds, table, col, desc, model, ts])
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        wb.save(tmp)
        os.replace(tmp, path)
        return True
    except (OSError, ValueError, KeyError, zipfile.BadZipFile, InvalidFileException) as e:
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except OSError:
            pass
        logger.exception("descriptions workbook save failed: %s", path)
        notify("error", f"Could not write `{path}` ({e}) — if the file is open in "
               "Excel, close it and regenerate. The descriptions are still "
               "available in this session.")
        return False


def flatten_descriptions(descriptions: dict) -> list:
    """Preview rows for st.dataframe (table-level row first per table)."""
    return [{"table": table, "column": col, "description": desc}
            for table, info in descriptions.items()
            for col, desc in [("(table)", info["table_description"]),
                              *info["columns"].items()]]


def null_description_labels(descriptions: dict, tagged: dict) -> tuple[list, list]:
    """(failed, policy_excluded) 'table' / 'table.column' labels for the null
    descriptions, split by cause for the post-generation report."""
    failed, excluded = [], []
    for table, info in descriptions.items():
        if info.get("table_description") is None:
            failed.append(table)
        for col, desc in info.get("columns", {}).items():
            if desc is None:
                (excluded if _is_tagged(col, tagged.get(table, ()))
                 else failed).append(f"{table}.{col}")
    return failed, excluded


def active_descriptions(descriptions, enabled: bool):
    """Descriptions for prompt injection — None while the feature is off, so a
    stale session/Excel can never leak into prompts; null descriptions pruned,
    so neither nulls nor policy-tagged names reach the action-plan/rule prompts."""
    if not enabled:
        return None
    pruned = {}
    for table, info in (descriptions or {}).items():
        entry = {k: v for k, v in (
            ("table_description", info.get("table_description")),
            ("columns", {c: d for c, d in info.get("columns", {}).items() if d})) if v}
        if entry:
            pruned[table] = entry
    return pruned or None


# --- YAML rendering (dataplex-dq scan format) -------------------------------------
_I_DASH, _I_KEY, _I_SUB = " " * 14, " " * 16, " " * 18
_EXPECTATIONS = {
    "non_null_expectation", "uniqueness_expectation", "range_expectation",
    "set_expectation", "regex_expectation", "table_condition_expectation",
}


def _num(value) -> str:
    """Bare number when finite (whole floats without decimal), else a quoted
    scalar — so string bounds like '2020-01-01' can't corrupt the document."""
    n = finite_float(value)
    if n is None:
        return yaml_quote(value)
    return str(int(n)) if n == int(n) else str(n)


def rule_is_valid(rule) -> bool:
    """Renderable = name + dimension + a KNOWN expectation type; unknown types
    would otherwise render as an expectation-less rule Dataplex rejects."""
    if not (isinstance(rule, dict) and bool(rule.get("name"))
            and bool(rule.get("dimension"))
            and (rule.get("expectation") or {}).get("type") in _EXPECTATIONS):
        return False
        
    exp = rule.get("expectation", {})
    etype = exp.get("type")
    column = str(rule.get("column", "")).lower()
    
    # GUARDRAIL 1: Prevent set/range rules on high-cardinality ID/Name columns
    if etype in ("set_expectation", "range_expectation"):
        if column.endswith(("_id", "_pin", "_num", "_name", "_nm")):
            return False
            
    # GUARDRAIL 2: Prevent hardcoded row counts in table conditions
    if etype == "table_condition_expectation":
        sql = str(exp.get("sql_expression", "")).upper()
        # Allow COUNT(*) > 0, but block exact bounds like BETWEEN or = 
        if "COUNT(*)" in sql and ("BETWEEN" in sql or "=" in sql):
            return False

    return True


def _sanitize_rule(rule):
    """Coerce one LLM rule toward validity: Dataplex name charset, float
    threshold (dropped if not coercible), stripped column/dimension, list-typed
    set values. Intercepts and corrects known stubborn LLM hallucinations."""
    if not isinstance(rule, dict):
        return rule
    if rule.get("name"):
        rule["name"] = re.sub(r"[^a-z0-9-]+", "-", str(rule["name"]).lower()).strip("-")
    if rule.get("column") is not None:
        rule["column"] = str(rule["column"]).strip()
    if rule.get("dimension") is not None:
        rule["dimension"] = str(rule["dimension"]).strip().upper()
    if rule.get("threshold") is not None:
        rule["threshold"] = finite_float(str(rule["threshold"]).strip().rstrip("%"))
        if rule["threshold"] is None:
            del rule["threshold"]
            
    exp = rule.get("expectation")
    if isinstance(exp, dict):
        if exp.get("type") == "set_expectation":
            values = exp.get("values")
            exp["values"] = values if isinstance(values, list) else ([] if values is None else [values])
            
        elif exp.get("type") == "regex_expectation":
            # GUARDRAIL 3: Fix phone number regex if the LLM falls back to strict 10 digits
            col = str(rule.get("column", "")).lower()
            if "phone" in col and exp.get("regex") == "^[0-9]{10}$":
                exp["regex"] = "^[0-9]{10,15}$"
                
    return rule


def render_rule(rule: dict) -> list:
    exp = rule.get("expectation") or {}
    etype, column = exp.get("type"), rule.get("column")
    lines = ([_I_DASH + f"- column: {yaml_quote(column)}",
              _I_KEY + f"name: {yaml_quote(rule['name'])}"]
             if column else [_I_DASH + f"- name: {yaml_quote(rule['name'])}"])
    lines.append(_I_KEY + f"dimension: {yaml_quote(rule['dimension'])}")
    threshold = finite_float(rule.get("threshold"))
    if etype != "table_condition_expectation" and threshold is not None:
        lines.append(_I_KEY + f"threshold: {threshold}")
    if etype in ("non_null_expectation", "uniqueness_expectation"):
        lines.append(_I_KEY + f"{etype}: true")
    elif etype == "range_expectation":
        lines.append(_I_KEY + "range_expectation:")
        for key in ("min_value", "max_value"):
            if exp.get(key) is not None:
                lines.append(_I_SUB + f"{key}: {_num(exp[key])}")
    elif etype == "set_expectation":
        lines.append(_I_KEY + "set_expectation:")
        lines.append(_I_SUB + "values: ["
                     + ", ".join(yaml_quote(v) for v in exp.get("values") or []) + "]")
        if exp.get("ignore_null"):
            lines.append(_I_SUB + "ignore_null: true")
    elif etype == "regex_expectation":
        lines.append(_I_KEY + "regex_expectation:")
        lines.append(_I_SUB + f"regex: {yaml_quote(exp.get('regex', ''))}")
        if exp.get("ignore_null"):
            lines.append(_I_SUB + "ignore_null: true")
    else:  # table_condition_expectation
        lines.append(_I_KEY + "table_condition_expectation:")
        lines.append(_I_SUB + f"sql_expression: {yaml_quote(exp.get('sql_expression', ''))}")
    return lines


def render_scan_block(scan_id, project_id_, dataset_id_, table_id_, rules) -> str:
    lines = [
        f"        {scan_id}:",
        "          data_source:",
        f"            project_id: {yaml_quote(project_id_)}",
        f"            dataset_id: {yaml_quote(dataset_id_)}",
        f"            table_id: {yaml_quote(table_id_)}",
        "          data_quality_spec:",
        "            rules:",
    ]
    for rule in rules:
        lines.extend(render_rule(rule))
    return "\n".join(lines)


# --- LLM steps ---------------------------------------------------------------------
_SYS_PLAN = """You are an expert Google Cloud Dataplex data quality engineer.
From the provided BigQuery column profiling data, draft a step-by-step Action Plan for a Dataplex DQ scan:
1. Identify columns suited to non_null_expectation (COMPLETENESS), uniqueness_expectation (UNIQUENESS), range/set/regex_expectation (VALIDITY), or table-level table_condition_expectation (VOLUME/FRESHNESS).
2. Justify each suggestion from the statistics (e.g. "range_expectation for antenna_face: values span 1-3 with no outliers"; "uniqueness_expectation for acct_id: percent_unique is 100%").
3. DO NOT suggest `set_expectation` or `range_expectation` for high-cardinality identifiers (e.g., columns ending in `_id`, `_pin`, `_name`, or general free-text fields). Only use `set_expectation` for known low-cardinality enums or codes. DO NOT suggest hardcoded `table_condition_expectation` row counts based on current sample sizes.
4. Do NOT write YAML/JSON yet — analysis and justification only, in clear markdown."""

_SYS_RULES = """You are an expert Google Cloud Dataplex data quality engineer.
From column profiling data, an action plan and user feedback, return ONLY one JSON object (no markdown, no prose) of this shape:
{"tables": {"<table_name>": {"rules": [{
  "column": "<column_name>",   // omit for table-level rules
  "name": "<rule-name>",       // lowercase letters, digits and hyphens ONLY
  "dimension": "COMPLETENESS|UNIQUENESS|VALIDITY|VOLUME|FRESHNESS",
  "threshold": 0.99,           // float; omit for table-level rules
  "expectation": {...}}]}}}    // exactly one of:
- {"type":"non_null_expectation"}                                        COMPLETENESS, threshold 0.99
- {"type":"uniqueness_expectation"}                                      UNIQUENESS, threshold 1.0
- {"type":"range_expectation","min_value":0,"max_value":9}               VALIDITY, threshold 0.99 (min and/or max)
- {"type":"set_expectation","values":["A","B"],"ignore_null":true}       VALIDITY, threshold 0.99
- {"type":"regex_expectation","regex":"^[0-9]{10,15}$","ignore_null":true} VALIDITY, threshold 0.99
- {"type":"table_condition_expectation","sql_expression":"COUNT(*) > 0"} VOLUME/FRESHNESS; NO column, NO threshold

CRITICAL INSTRUCTIONS:
- DO NOT apply `set_expectation` to high-cardinality entity IDs (e.g., `kb_sales_rep_pin`, `outlet_id`, `servreq_header_nm`, `sales_rep_id`, `operator_id`, `kb_dealer_cd`, `chnl_org_id`). 
- DO NOT set hardcoded, restrictive range boundaries for sequential keys (like `servreq_header_id`) or overall table volume (avoid strictly bounding `COUNT(*)`).
- When validating phone numbers via regex, allow reasonable variations in length (e.g., `^[0-9]{10,15}$`).
Use table_name keys exactly as given. Only reference columns present in the profiling data. Incorporate all user feedback."""


_DESC_LABEL = ("Table & Column Business Descriptions (use them to infer "
               "business rules the statistics alone cannot show):\n")

_PROMPT_CHAR_CAP = 400_000  # generous whole-prompt budget for plan/rules calls
_UPLOAD_CHAR_CAP = 60_000   # reference-file share of the rules prompt


def _slim_profiles(profile_json):
    """Profile rows without their top_n detail — the lossy lever when a
    plan/rules prompt exceeds the budget."""
    return [{k: v for k, v in p.items() if k != "top_n"} for p in profile_json]


def generate_action_plan(profile_json, *, call_llm, notify, descriptions=None):
    def build(profiles):
        prompt = ("Analyze this column profiling data and draft a data quality action plan:\n"
                  + _jsonc(profiles))
        if descriptions:
            prompt += "\n\n" + _DESC_LABEL + _jsonc(descriptions)
        return prompt

    prompt = build(profile_json)
    if len(prompt) > _PROMPT_CHAR_CAP:
        prompt = build(_slim_profiles(profile_json))
        notify("warning", "Profiling statistics exceeded the prompt budget — "
               "top-value details were dropped from the action-plan prompt.")
    return call_llm(_SYS_PLAN, prompt)


def _read_upload(uploaded_file, *, notify) -> str:
    """Best-effort text of an uploaded reference file; '' if binary."""
    if uploaded_file is None:
        return ""
    raw = uploaded_file.read()
    for enc in ("utf-8", "latin-1"):
        try:
            text = raw.decode(enc)
            if text.strip():
                return text
        except (UnicodeDecodeError, AttributeError):
            pass
    notify("warning", f"'{getattr(uploaded_file, 'name', 'file')}' is not readable as text "
           "(binary formats are unsupported) — paste key requirements into the feedback box.")
    return ""


def generate_rules(profile_json, action_plan, hitl_feedback, *, call_llm, notify,
                   uploaded_file=None, descriptions=None):
    """{table_name: [rule_dict, ...]} from the model; sanitized per rule."""
    desc_section = _DESC_LABEL + _jsonc(descriptions) + "\n\n" if descriptions else ""
    reference = _read_upload(uploaded_file, notify=notify) or "(none provided)"
    if len(reference) > _UPLOAD_CHAR_CAP:
        reference = reference[:_UPLOAD_CHAR_CAP]
        notify("warning", f"Reference document truncated to its first "
               f"{_UPLOAD_CHAR_CAP:,} characters for the prompt.")

    def build(profiles):
        return (f"Profiling Data:\n{_jsonc(profiles)}\n\n"
                + desc_section +
                f"Proposed Action Plan:\n{action_plan}\n\n"
                f"User Adjustments/HITL Feedback:\n{hitl_feedback}\n\n"
                f"Reference document contents (if any):\n{reference}\n\n"
                "Return the JSON object of rules per table.")

    prompt = build(profile_json)
    if len(prompt) > _PROMPT_CHAR_CAP:
        prompt = build(_slim_profiles(profile_json))
        notify("warning", "Profiling statistics exceeded the prompt budget — "
               "top-value details were dropped from the rules prompt.")
    raw_text = call_llm(_SYS_RULES, prompt)
    data = parse_llm_json(raw_text)
    if data is None:
        raise ValueError("Model did not return JSON. Raw response:\n" + raw_text)
    tables = data.get("tables", data) if isinstance(data, dict) else {}
    return {tbl: [_sanitize_rule(r) for r in
                  (spec.get("rules", []) if isinstance(spec, dict)
                   else spec if isinstance(spec, list) else [])]
            for tbl, spec in tables.items()}


def build_scan_plan(table_order, rules_by_table, settings: Settings):
    """(plan, blocks, scan_ids): report rows plus rendered blocks and their ids
    for tables that passed scan-id validation and have valid rules."""
    forbidden = collect_repo_scan_ids(settings.governance_dir, settings.instance)
    plan, blocks, ids = [], [], []
    for table in table_order:
        rules = [r for r in (rules_by_table.get(table) or []) if rule_is_valid(r)]
        scan_id = resolve_scan_id(settings.instance, settings.source_dataset_id,
                                  table, forbidden | set(ids))
        ok, reason = validate_scan_id(scan_id, settings.instance, forbidden | set(ids))
        if ok and not rules:
            ok, reason = False, "no valid rules generated"
        if ok:
            blocks.append(render_scan_block(scan_id, settings.source_project_id,
                                            settings.source_dataset_id, table, rules))
            ids.append(scan_id)
        plan.append({"table": table, "scan_id": scan_id, "rules": len(rules),
                     "status": "ok" if ok else f"skip: {reason}"})
    return plan, blocks, ids