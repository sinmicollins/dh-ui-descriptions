"""dq_generation: prompt byte-identity, the policy-tag firewall, description
generation (sentinel nulls, hybrid mode, Collibra enrichment, context budget),
the Excel store, and the parameterized sampling SQL cascade.

Ported 1:1 from the pre-refactor behavior suite — every monkey-patched module
global became an injected parameter."""
import json
import os
from datetime import datetime, timezone
from types import SimpleNamespace

import openpyxl
import pytest
from google.api_core.exceptions import BadRequest
from google.cloud.bigquery import SchemaField
from google.cloud.bigquery.schema import PolicyTagList

import dq_generation as gen
from conftest import FakeResp, Notes, blank_meta, make_settings

PROFILES = [{"table_name": "t1", "column_name": "c1", "column_type": "STRING",
             "percent_null": 0.5, "top_n": [{"value": "A", "count": 9, "percent": 90.0}]}]
DESCS = {"t1": {"table_description": "This table contains activations.",
                "columns": {"c1": "This column contains a code."}}}
SETTINGS = make_settings()


def capture_llm(prompts, reply='{"columns": {}}'):
    def call_llm(system, user):
        prompts.append(user)
        return reply
    return call_llm


def rows_stub(rows):
    return lambda table, meta, exclude: [dict(r) for r in rows]


# --- 1. off-mode privacy contract: prompts byte-identical to the old format --------
def test_offmode_prompts_byte_identical(notes):
    captured = {}

    def call_llm(system, user):
        captured["system"], captured["user"] = system, user
        return '{"tables": {}}'

    gen.generate_action_plan(PROFILES, call_llm=call_llm, notify=notes)
    old_plan_prompt = ("Analyze this column profiling data and draft a data quality action plan:\n"
                       + json.dumps(PROFILES, separators=(",", ":"), default=str))
    assert captured["user"] == old_plan_prompt, "off-mode action-plan prompt changed!"

    gen.generate_rules(PROFILES, "the plan", "looks good", call_llm=call_llm, notify=notes)
    old_rules_prompt = (f"Profiling Data:\n{json.dumps(PROFILES, separators=(',', ':'), default=str)}\n\n"
                        f"Proposed Action Plan:\nthe plan\n\n"
                        f"User Adjustments/HITL Feedback:\nlooks good\n\n"
                        f"Reference document contents (if any):\n(none provided)\n\n"
                        "Return the JSON object of rules per table.")
    assert captured["user"] == old_rules_prompt, "off-mode rules prompt changed!"


# --- 2. on-mode: descriptions injected into plan/rules prompts -----------------------
def test_descriptions_injected_when_enabled(notes):
    captured = {}

    def call_llm(system, user):
        captured["user"] = user
        return '{"tables": {}}'

    gen.generate_action_plan(PROFILES, call_llm=call_llm, notify=notes, descriptions=DESCS)
    assert "Table & Column Business Descriptions" in captured["user"]
    assert "This table contains activations." in captured["user"]
    gen.generate_rules(PROFILES, "the plan", "looks good", call_llm=call_llm,
                       notify=notes, descriptions=DESCS)
    assert "This column contains a code." in captured["user"]


# --- 3. active_descriptions honors the toggle ----------------------------------------
def test_active_descriptions_gate():
    assert gen.active_descriptions(DESCS, enabled=False) is None, \
        "descriptions leaked while toggle off!"
    assert gen.active_descriptions(DESCS, enabled=True) == DESCS


# --- 4. build_sample_evidence ---------------------------------------------------------
def test_build_sample_evidence():
    rows = [{"code": "A" * 200, "ts": f"2026-01-{i % 9 + 1:02d}",
             "maybe": None if i % 2 else i} for i in range(50)]
    ev = gen.build_sample_evidence(rows, max_rows=5, top_values=3, cell_cap=10)
    assert set(ev) == {"columns", "recent_rows"}
    assert len(ev["recent_rows"]) == 5
    assert ev["columns"]["code"]["top_values"][0]["value"] == "A" * 10 + "..."
    assert ev["columns"]["code"]["top_values"][0]["count"] == 50
    assert ev["columns"]["maybe"]["non_null"] == 25
    assert ev["columns"]["maybe"]["rows_sampled"] == 50
    assert len(ev["columns"]["ts"]["top_values"]) == 3
    assert gen.build_sample_evidence([]) == {}


# --- 5. Excel upsert -------------------------------------------------------------------
def save(path, descriptions, notes):
    return gen.save_descriptions_xlsx(str(path), descriptions, dataset="ent_actvn",
                                      model="model", notify=notes)


def test_xlsx_upsert(tmp_path, notes):
    xlsx = tmp_path / "test_descs.xlsx"
    assert save(xlsx, DESCS, notes)
    descs2 = {"t1": {"table_description": "This table contains UPDATED activations.",
                     "columns": {"c1": "This column contains a code.",
                                 "c2": "This column contains a date."}},
              "t2": {"table_description": "This table contains payments.", "columns": {}}}
    assert save(xlsx, descs2, notes)
    ws = openpyxl.load_workbook(xlsx).active
    data = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    keys = [(r[0], r[1], r[2] or "") for r in data]
    assert len(keys) == len(set(keys)), f"duplicate keys after upsert: {keys}"
    by_key = {(r[0], r[1], r[2] or ""): r[3] for r in data}
    assert by_key[("ent_actvn", "t1", "")] == "This table contains UPDATED activations."
    assert by_key[("ent_actvn", "t1", "c2")] == "This column contains a date."
    assert by_key[("ent_actvn", "t2", "")] == "This table contains payments."
    assert [c.value for c in ws[1]] == gen._XLSX_HEADER


def test_xlsx_hardened_save(tmp_path, notes):
    """Sheet targeted by name, whitespace-key duplicates synced, illegal chars
    stripped, unrelated sheets untouched."""
    xlsx2 = tmp_path / "test_descs2.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "descriptions"
    ws.append(gen._XLSX_HEADER)
    ws.append(["ent_actvn", "t1", "", "stale one", "m", "old"])
    ws.append(["ent_actvn", "t1 ", "", "stale duplicate (whitespace key)", "m", "old"])
    other = wb.create_sheet("notes")
    other.append(["unrelated"])
    wb.active = other  # active sheet is NOT the descriptions sheet
    wb.save(xlsx2)

    descs3 = {"t1": {"table_description": "This table contains fresh text.\x00\x01",
                     "columns": {}}}
    assert save(xlsx2, descs3, notes)
    wb = openpyxl.load_workbook(xlsx2)
    assert wb["notes"]["A1"].value == "unrelated"          # other sheet untouched
    rows = [[c.value for c in r] for r in wb["descriptions"].iter_rows(min_row=2)]
    assert len(rows) == 2, rows                            # no third row appended
    assert all(r[3] == "This table contains fresh text." for r in rows), rows


def test_xlsx_refusals(tmp_path, notes):
    """Header mismatch refused; locked file returns a friendly False."""
    descs3 = {"t1": {"table_description": "This table contains fresh text.",
                     "columns": {}}}
    xlsx3 = tmp_path / "test_descs3.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "descriptions"
    ws.append(["colA", "colB"])  # wrong header
    wb.save(xlsx3)
    assert not save(xlsx3, descs3, notes)
    assert "unexpected" in notes.errors[-1], notes.errors

    xlsx2 = tmp_path / "test_descs2.xlsx"
    assert save(xlsx2, descs3, notes)
    with open(xlsx2, "rb"):
        locked_result = save(xlsx2, descs3, notes)
    assert locked_result is False and "Could not write" in notes.errors[-1], notes.errors
    assert not os.path.exists(str(xlsx2) + ".tmp"), "tmp file left behind"


# --- 6. flatten + path shape --------------------------------------------------------
def test_flatten_and_xlsx_path():
    descs2 = {"t1": {"table_description": "This table contains UPDATED activations.",
                     "columns": {"c1": "x"}}}
    flat = gen.flatten_descriptions(descs2)
    assert flat[0] == {"table": "t1", "column": "(table)",
                       "description": "This table contains UPDATED activations."}
    path = gen.descriptions_xlsx_path(make_settings(instance="dh1",
                                                    source_dataset_id="ent_actvn"))
    assert path.endswith(os.path.join("descriptions", "dh1_dq_descriptions_ent_actvn.xlsx"))


# --- 7. Collibra glossary enrichment --------------------------------------------------
GLOSSARY_BLOCK = "TELUS BUSINESS GLOSSARY"


class CollibraHttp:
    """Bearer-accepting Collibra fake: probe, two glossary pages, definitions."""

    def __init__(self):
        self.calls = []

    def get(self, url, params=None, headers=None, timeout=None, **kw):
        params = dict(params or {})
        if url.endswith("/rest/2.0/assets") and "offset" not in params:
            return FakeResp({"total": 0, "results": []})  # auth-scheme probe
        self.calls.append((url, params))
        assert headers["Authorization"] == "Bearer KEY"
        if url.endswith("/rest/2.0/assets"):
            assert params["typeIds"] == gen._BUSINESS_TERM_TYPE_ID
            assert "domainId" not in params  # blank domain omitted
            if params["offset"] == 0:
                return FakeResp({"total": 3, "results": [
                    {"id": "id-5g", "name": "5G", "displayName": "5th Generation"},
                    {"id": "id-acct", "name": "ACCT", "displayName": "Account"}]})
            return FakeResp({"total": 3, "results": [
                {"id": "id-actvn", "name": "ACTVN", "displayName": "Activation"}]})
        if url.endswith("/rest/2.0/attributes"):
            return FakeResp({"results": [{"type": {"name": "Definition"},
                                          "value": "<p>The fifth generation  of mobile networks.</p>"}]})
        raise AssertionError("unexpected URL " + url)

    def post(self, *a, **k):
        raise AssertionError("no session login expected for a colon-free key")


def collibra_settings(**over):
    return make_settings(use_collibra=True, collibra_url="https://x.collibra.com",
                         collibra_domain="", **over)


def run_descs(settings, notes, prompts, tables=("bq_actvn_txn",), *,
              load_glossary=None, reply='{"columns": {}}', abbrev={}, **kw):
    return gen.generate_descriptions(
        list(tables), [], kw.pop("tagged", {}), settings,
        call_llm=kw.pop("call_llm", capture_llm(prompts, reply)),
        fetch_table_meta=kw.pop("fetch_table_meta", lambda *a: blank_meta(*tables)),
        fetch_rows=kw.pop("fetch_rows", rows_stub([{"acct_5g_ind": "Y", "other_cd": "1"}])),
        notify=notes, load_glossary=load_glossary, abbrev=abbrev, **kw)


def test_collibra_import_paging_matching_injection(notes):
    http = CollibraHttp()
    prompts = []
    settings = collibra_settings()
    load_glossary = lambda: gen.load_collibra(settings, get_key=lambda: "KEY", http=http)
    run_descs(settings, notes, prompts, load_glossary=load_glossary)
    assert all(GLOSSARY_BLOCK in p for p in prompts), "glossary block missing from prompts"
    assert "5G = 5th Generation — The fifth generation of mobile networks." in prompts[0]
    assert "ACTVN = Activation" in prompts[0]  # matched from the table name token
    assert "ACCT = Account" in prompts[0]      # matched from column acct_5g_ind
    asset_calls = [c for c in http.calls if c[0].endswith("/assets")]
    attr_calls = [c for c in http.calls if c[0].endswith("/attributes")]
    assert len(asset_calls) == 2, f"paging expected 2 asset calls, got {len(asset_calls)}"
    assert {c[1]["assetId"] for c in attr_calls} == {"id-5g", "id-acct", "id-actvn"}


# --- 8. Collibra off / failure paths ---------------------------------------------------
def test_collibra_disabled_and_failure_degrade(notes):
    prompts, loads = [], []
    run_descs(make_settings(), notes, prompts,
              load_glossary=lambda: loads.append(1) or ({}, None))
    assert not loads, "Collibra was called while disabled!"
    assert not any(GLOSSARY_BLOCK in p for p in prompts)

    prompts.clear()
    settings = collibra_settings()

    def failing_load():
        return gen.load_collibra(
            settings, get_key=lambda: (_ for _ in ()).throw(RuntimeError("no ADC")),
            http=SimpleNamespace())

    out = run_descs(settings, notes, prompts, load_glossary=failing_load)
    assert out and notes.warnings and "Collibra glossary unavailable" in notes.warnings[0]
    assert not any(GLOSSARY_BLOCK in p for p in prompts)


# --- 9. auth cascade: Bearer rejected, Basic (user:pass) accepted ------------------------
class CascadeHttp:
    def __init__(self):
        self.calls = []

    def get(self, url, params=None, headers=None, auth=None, cookies=None, timeout=None):
        params = dict(params or {})
        self.calls.append((url, params, headers, auth))
        if url.endswith("/rest/2.0/assets"):
            if headers and "Authorization" in headers:
                return FakeResp({}, 401)          # Bearer and pre-encoded Basic rejected
            if auth == ("svc_user", "s3cret"):
                if "offset" not in params:
                    return FakeResp({"total": 0, "results": []})  # probe OK
                return FakeResp({"total": 1, "results": [
                    {"id": "id-acct", "name": "ACCT", "displayName": "Account"}]})
            return FakeResp({}, 401)
        if url.endswith("/rest/2.0/attributes"):
            assert auth == ("svc_user", "s3cret")
            return FakeResp({"results": []})
        raise AssertionError(url)

    def post(self, *a, **k):
        return FakeResp({}, 401)


def test_collibra_auth_cascade(notes):
    http = CascadeHttp()
    prompts = []
    settings = collibra_settings()
    load_glossary = lambda: gen.load_collibra(settings, get_key=lambda: "svc_user:s3cret",
                                              http=http)
    run_descs(settings, notes, prompts, load_glossary=load_glossary)
    assert not notes.warnings, notes.warnings
    assert any("ACCT = Account" in p for p in prompts), "Basic-auth glossary not injected"
    paged = [c for c in http.calls if c[0].endswith("/assets") and "offset" in c[1]]
    assert all(c[3] == ("svc_user", "s3cret") and not c[2] for c in paged)

    # all schemes rejected -> one warning, generation continues
    reject_all = SimpleNamespace(get=lambda *a, **k: FakeResp({}, 401),
                                 post=lambda *a, **k: FakeResp({}, 401))
    prompts.clear()
    fail_notes = Notes()
    load_glossary = lambda: gen.load_collibra(settings, get_key=lambda: "svc_user:s3cret",
                                              http=reject_all)
    out = run_descs(settings, fail_notes, prompts, load_glossary=load_glossary)
    assert out and fail_notes.warnings and \
        "rejected every auth scheme" in fail_notes.warnings[0], fail_notes.warnings


# --- 9b. saved glossary CSV: round trip, refresh merge, offline reuse ---------------------
def glossary_of(*triples):
    return {name.lower(): {"id": aid, "name": name, "full": full}
            for name, full, aid in triples}


def test_collibra_csv_round_trip(tmp_path):
    path = str(tmp_path / "collibra_glossary.csv")
    glossary = glossary_of(("5G", "5th Generation", "id-5g"),
                           ("ACCT", "Compte, señor", "id-acct"),
                           ("ACTVN", "Activation", "id-actvn"))
    gen.save_collibra_csv(path, glossary, {"id-5g": "Fifth generation, of networks"})
    loaded, fetch_defs = gen.load_collibra_csv(path)
    assert loaded == glossary
    assert fetch_defs(("id-5g", "id-acct")) == {"id-5g": "Fifth generation, of networks"}


def test_collibra_csv_refresh_preserves_surviving_definitions(tmp_path):
    path = str(tmp_path / "collibra_glossary.csv")
    gen.save_collibra_csv(path, glossary_of(("A", "Aaa", "id-a"), ("B", "Bbb", "id-b")),
                          {"id-a": "old A def", "id-b": "old B def"})
    gen.save_collibra_csv(path, glossary_of(("A", "Aaa", "id-a"), ("C", "Ccc", "id-c")),
                          {"id-c": "new C def"})
    loaded, fetch_defs = gen.load_collibra_csv(path)
    assert set(loaded) == {"a", "c"}
    assert fetch_defs(("id-a", "id-b", "id-c")) == {"id-a": "old A def",
                                                    "id-c": "new C def"}


def test_load_collibra_persists_retrieval_then_offline_reuse(tmp_path, notes):
    path = str(tmp_path / "collibra_glossary.csv")
    settings = collibra_settings()
    prompts = []
    load_glossary = lambda: gen.load_collibra(settings, get_key=lambda: "KEY",
                                              http=CollibraHttp(), save_path=path)
    run_descs(settings, notes, prompts, load_glossary=load_glossary)
    loaded, fetch_defs = gen.load_collibra_csv(path)
    assert set(loaded) == {"5g", "acct", "actvn"}
    assert fetch_defs(("id-5g",)) == {"id-5g": "The fifth generation of mobile networks."}

    # a later run off the saved CSV injects the same hints with zero Collibra traffic
    prompts.clear()
    offline_notes = Notes()
    run_descs(settings, offline_notes, prompts,
              load_glossary=lambda: gen.load_collibra_csv(path))
    assert "5G = 5th Generation — The fifth generation of mobile networks." in prompts[0]
    assert not offline_notes.warnings, offline_notes.warnings


def test_load_collibra_csv_missing_degrades(tmp_path, notes):
    missing = str(tmp_path / "nope.csv")
    with pytest.raises(RuntimeError, match="saved Collibra glossary unreadable"):
        gen.load_collibra_csv(missing)
    prompts = []
    out = run_descs(collibra_settings(), notes, prompts,
                    load_glossary=lambda: gen.load_collibra_csv(missing))
    assert out and notes.warnings and "Collibra glossary unavailable" in notes.warnings[0]
    assert not any(GLOSSARY_BLOCK in p for p in prompts)


# --- 10. abbreviations.csv hints ----------------------------------------------------------
def test_abbreviation_hints(notes):
    prompts = []
    run_descs(make_settings(), notes, prompts,
              abbrev={"account": "acct", "activation": "actvn", "bq_": ""})
    assert "TELUS ABBREVIATION GLOSSARY" in prompts[0]
    assert "acct = account" in prompts[0] and "actvn = activation" in prompts[0]
    assert "bq =" not in prompts[0]


# --- 11. per-table failure isolation --------------------------------------------------------
def test_per_table_failure_isolation(notes):
    prompts = []

    def picky_llm(system, user):
        prompts.append(user)
        if "TABLE NAME: t_bad" in user:
            raise RuntimeError("FuelIX 400: request too large")
        return ("This table contains good data." if system is gen._SYS_TABLE_DESC
                else '{"columns": {}}')

    out = run_descs(make_settings(), notes, prompts, tables=("t_bad", "t_good"),
                    call_llm=picky_llm)
    assert out["t_bad"]["table_description"] is None
    assert out["t_bad"]["columns"] == {"acct_5g_ind": None, "other_cd": None}
    assert out["t_good"]["table_description"] == "This table contains good data."
    assert len(notes.warnings) == 1 and "t_bad" in notes.warnings[0] \
        and "400" in notes.warnings[0]


# --- 12. context size cap on wide tables ------------------------------------------------------
WIDE_ROW = {f"column_{i}_padpadpad": "v" * 200 for i in range(300)}


def test_context_capped_on_wide_tables(notes):
    prompts = []
    run_descs(make_settings(), notes, prompts, tables=("wide",),
              fetch_rows=rows_stub([WIDE_ROW] * 30))
    assert prompts and all(len(p) < gen._CONTEXT_CHAR_CAP + 10_000 for p in prompts), \
        [len(p) for p in prompts]


def test_context_budget_includes_reference_and_glossaries(notes):
    """The 80k cap now covers the appended blocks too (it used to be bypassed)."""
    prompts = []
    existing = {"wide": {"table_description": "Existing table text. " * 2000,
                         "columns": {}}}
    run_descs(make_settings(), notes, prompts, tables=("wide",),
              fetch_rows=rows_stub([WIDE_ROW] * 30), existing=existing)
    table_prompt = prompts[0]
    assert "EXISTING BIGQUERY DESCRIPTIONS" in table_prompt
    assert len(table_prompt) <= gen._CONTEXT_CHAR_CAP + 100, len(table_prompt)


# --- 14. null instead of the sentinel phrase ----------------------------------------------------
def test_sentinel_and_blank_stored_as_null(notes):
    assert gen._null_if_unable("Unable to Generate Description") is None
    assert gen._null_if_unable("  unable to generate description.  ") is None
    assert gen._null_if_unable("") is None and gen._null_if_unable(None) is None
    assert gen._null_if_unable("This table contains x.") == "This table contains x."

    def unable_llm(system, user):
        if system is gen._SYS_TABLE_DESC:
            return "Unable to Generate Description"
        return ('{"columns": {"col_a": "Unable to Generate Description.",'
                ' "col_b": "This column contains a code."}}')

    out = run_descs(make_settings(), notes, [], tables=("t_null",), call_llm=unable_llm,
                    fetch_rows=rows_stub([{"col_a": "1", "col_b": "2"}]))
    assert out["t_null"]["table_description"] is None
    assert out["t_null"]["columns"] == {"col_a": None,
                                        "col_b": "This column contains a code."}


# --- 15. fetch_bq_metadata + _is_tagged (real SchemaFields) --------------------------------------
SCHEMA = [
    SchemaField("plain", "STRING", description="The plain column."),
    SchemaField("secret_col", "STRING", description="SIN of the customer.",
                policy_tags=PolicyTagList(names=("projects/p/taxonomies/1/policyTags/2",))),
    SchemaField("rec", "RECORD", fields=(
        SchemaField("leaf", "STRING", policy_tags=PolicyTagList(names=("t",))),
        SchemaField("open", "STRING", description="Open leaf."))),
]


class PolicyBqClient:
    def __init__(self):
        self.sqls = []

    def get_table(self, ref):
        assert ref == "src-proj.src_ds.t1", ref
        return SimpleNamespace(schema=SCHEMA, description="Existing table text.")

    def query(self, sql, job_config=None):
        self.sqls.append(sql)
        return [] if "INFORMATION_SCHEMA" in sql else [{"plain": "v"}]


def policy_meta():
    return gen.fetch_bq_metadata(PolicyBqClient(), "src-proj", "src_ds", ("t1",))["t1"]


def test_fetch_bq_metadata_and_is_tagged():
    meta1 = policy_meta()
    assert meta1["tagged"] == ("secret_col", "rec.leaf"), meta1
    assert meta1["table_description"] == "Existing table text."
    assert meta1["columns"] == {"plain": "The plain column.",
                                "secret_col": "SIN of the customer.", "rec": None,
                                "rec.leaf": None, "rec.open": "Open leaf."}, meta1
    assert gen._is_tagged("secret_col", meta1["tagged"])
    assert gen._is_tagged("rec.leaf", meta1["tagged"])
    assert gen._is_tagged("rec.leaf.deeper", meta1["tagged"])   # nested under a tagged field
    assert gen._is_tagged("rec", meta1["tagged"])               # record containing a tagged leaf
    assert not gen._is_tagged("plain", meta1["tagged"])
    assert not gen._is_tagged("secret_col_2", meta1["tagged"])  # prefix but not a path segment


# --- 16. get_profiles_for_llm drops tagged rows + reports them ------------------------------------
def test_get_profiles_for_llm_policy_filter(notes):
    fetch_profiles = lambda *a: [
        {"table_name": "t1", "column_name": "plain", "percent_null": 0},
        {"table_name": "t1", "column_name": "secret_col", "percent_null": 0},
        {"table_name": "t1", "column_name": "rec.leaf", "percent_null": 0},
        {"table_name": "t2", "column_name": "free", "percent_null": 0},
    ]
    fetch_bq_meta = lambda *a: {
        "t1": {"tagged": ("secret_col", "rec.leaf"), "table_description": None, "columns": {}},
        "t2": {"tagged": (), "table_description": None, "columns": {}}}
    kept, tagged = gen.get_profiles_for_llm(["t1", "t2"], SETTINGS,
                                            fetch_profiles=fetch_profiles,
                                            fetch_bq_meta=fetch_bq_meta, notify=notes)
    assert [(p["table_name"], p["column_name"]) for p in kept] == \
        [("t1", "plain"), ("t2", "free")]
    assert notes.infos and "t1.secret_col" in notes.infos[0] \
        and "t1.rec.leaf" in notes.infos[0], notes.infos
    assert "t2" not in notes.infos[0].split("payload")[1].replace("t1.", ""), notes.infos


# --- 17. firewall inside generate_descriptions -----------------------------------------------------
def test_description_firewall_and_except_sql(notes):
    prompts, seen = [], {}

    def spy_rows(table, meta, exclude):
        seen["exclude"] = exclude
        return [{"open_col": "v"}]

    def desc_llm(system, user):
        prompts.append(user)
        if system is gen._SYS_TABLE_DESC:
            return "This table contains things."
        return '{"columns": {"open_col": "This column contains open data."}}'

    out = run_descs(make_settings(), notes, prompts, tables=("t_sec",),
                    call_llm=desc_llm, fetch_rows=spy_rows,
                    tagged={"t_sec": ("secret_col", "rec.leaf")})
    assert seen["exclude"] == {"secret_col", "rec"}, seen
    assert all("secret_col" not in p and "rec.leaf" not in p for p in prompts)
    assert out["t_sec"]["columns"] == {"open_col": "This column contains open data.",
                                       "secret_col": None, "rec.leaf": None}

    # real fetch_last_rows renders SELECT * EXCEPT so tagged data is never queried
    client = PolicyBqClient()
    meta0 = {"partition_column": "", "partition_column_type": "",
             "require_partition_filter": False, "temporal_columns": []}
    rows = gen.fetch_last_rows(client, "src-proj", "src_ds", "t1", meta0,
                               exclude=frozenset({"secret_col", "rec"}), notify=notes)
    assert rows == [{"plain": "v"}]
    assert client.sqls[-1].startswith(
        "SELECT * EXCEPT (`rec`, `secret_col`) FROM `src-proj.src_ds.t1`"), client.sqls


# --- 18. null_description_labels + _label_list ------------------------------------------------------
def test_null_labels_and_label_cap():
    descs_n = {"tA": {"table_description": None,
                      "columns": {"good": "text", "bad": None, "sec": None}},
               "tB": {"table_description": "ok", "columns": {}}}
    failed, excluded = gen.null_description_labels(descs_n, {"tA": ("sec",)})
    assert failed == ["tA", "tA.bad"] and excluded == ["tA.sec"], (failed, excluded)
    assert gen._label_list(["a", "b"]) == "`a`, `b`"
    capped = gen._label_list([f"c{i}" for i in range(65)])
    assert "and 5 more" in capped and capped.count("`") == 120


# --- 19. active_descriptions prunes nulls (and tagged names) from prompts ---------------------------
def test_active_descriptions_null_pruning():
    act = gen.active_descriptions(
        {"tA": {"table_description": None, "columns": {"good": "desc", "sec": None}},
         "tB": {"table_description": None, "columns": {"x": None}}}, enabled=True)
    assert act == {"tA": {"columns": {"good": "desc"}}}, act
    assert "sec" not in json.dumps(act) and "null" not in json.dumps(act)
    assert gen.active_descriptions(
        {"tB": {"table_description": None, "columns": {"x": None}}}, enabled=True) is None
    assert gen.active_descriptions(DESCS, enabled=False) is None  # off-mode gate unchanged


# --- 20. Excel: null -> blank cell, placeholder text cleared on upsert -------------------------------
def test_xlsx_null_cells_and_placeholder_cleared(tmp_path, notes):
    xlsx4 = tmp_path / "test_descs4.xlsx"
    assert save(xlsx4, {"t9": {
        "table_description": "Unable to Generate Description",   # legacy placeholder
        "columns": {"c1": "Unable to Generate Description"}}}, notes)
    assert save(xlsx4, {"t9": {
        "table_description": None,
        "columns": {"c1": None, "c2": "This column contains a code."}}}, notes)
    wb = openpyxl.load_workbook(xlsx4)
    rows = [[c.value for c in r] for r in wb["descriptions"].iter_rows(min_row=2)]
    by_key = {(r[0], r[1], r[2] or ""): r[3] for r in rows}
    assert by_key[("ent_actvn", "t9", "")] is None
    assert by_key[("ent_actvn", "t9", "c1")] is None
    assert by_key[("ent_actvn", "t9", "c2")] == "This column contains a code."


# --- 21. existing_descriptions_from: descriptions kept, tagged columns nulled ------------------------
def test_existing_descriptions_from_policy_nulled():
    existing1 = gen.existing_descriptions_from({"t1": policy_meta()})
    assert existing1["t1"]["table_description"] == "Existing table text."
    assert existing1["t1"]["columns"]["plain"] == "The plain column."
    assert existing1["t1"]["columns"]["rec.open"] == "Open leaf."
    assert existing1["t1"]["columns"]["secret_col"] is None    # policy-nulled
    assert existing1["t1"]["columns"]["rec.leaf"] is None      # policy-nulled
    assert existing1["t1"]["columns"]["rec"] is None           # parent of tagged leaf


# --- 22. generate-new mode: existing descriptions injected as reference only -------------------------
def test_existing_descriptions_as_reference(notes):
    prompts = []

    def ref_llm(system, user):
        prompts.append(user)
        if system is gen._SYS_TABLE_DESC:
            return "This table contains regenerated things."
        return ('{"columns": {"col_a": "This column contains new a.",'
                ' "col_b": "This column contains b."}}')

    existing_ref = {"t_ref": {"table_description": "Existing table text.",
                              "columns": {"col_a": "Existing col_a text.", "col_b": None}}}
    out = run_descs(make_settings(), notes, prompts, tables=("t_ref",), call_llm=ref_llm,
                    fetch_rows=rows_stub([{"col_a": "1", "col_b": "2"}]),
                    existing=existing_ref)
    assert all("EXISTING BIGQUERY DESCRIPTIONS" in p for p in prompts), "reference block missing"
    assert all("Existing table text." in p for p in prompts)
    first_line = [p for p in prompts if p.startswith("COLUMNS TO DESCRIBE")][0].splitlines()[0]
    assert '"col_a"' in first_line and '"col_b"' in first_line  # full regeneration asks for all
    assert out["t_ref"]["table_description"] == "This table contains regenerated things."
    assert out["t_ref"]["columns"]["col_a"] == "This column contains new a."


# --- 23. hybrid: keep current, generate only the missing ---------------------------------------------
def test_hybrid_fill_only_missing(notes):
    def _boom(*a, **k):
        raise AssertionError("should not be called for a fully described table")

    existing_full = {"t_full": {"table_description": "Existing table text.",
                                "columns": {"col_a": "A.", "col_b": "B.", "sec": None}}}
    out = run_descs(make_settings(), notes, [], tables=("t_full",), call_llm=_boom,
                    fetch_rows=_boom, tagged={"t_full": ("sec",)},
                    existing=existing_full, fill_only_missing=True)
    assert out["t_full"] == existing_full["t_full"], out

    prompts = []

    def gap_llm(system, user):
        prompts.append(user)
        assert system is gen._SYS_COL_DESC, "table description must not be regenerated"
        return '{"columns": {"col_b": "This column contains b."}}'

    existing_part = {"t_part": {"table_description": "Existing table text.",
                                "columns": {"col_a": "Existing col_a text.", "col_b": None}}}
    out = run_descs(make_settings(), notes, prompts, tables=("t_part",), call_llm=gap_llm,
                    fetch_rows=rows_stub([{"col_a": "1", "col_b": "2"}]),
                    existing=existing_part, fill_only_missing=True)
    assert len(prompts) == 1, [p[:80] for p in prompts]
    first_line = prompts[0].splitlines()[0]
    assert first_line.startswith("COLUMNS TO DESCRIBE") and '"col_b"' in first_line
    assert '"col_a"' not in first_line  # already described -> not asked
    assert "TABLE DESCRIPTION: Existing table text." in prompts[0]
    assert out["t_part"]["table_description"] == "Existing table text."
    assert out["t_part"]["columns"] == {"col_b": "This column contains b.",
                                        "col_a": "Existing col_a text."}
    assert not notes.warnings, notes.warnings


# --- sampling SQL: attempt cascade + typed parameters -------------------------------------------------
class CascadeClient:
    """PARTITIONS lookups succeed; every sample query fails with a BQ error."""

    def __init__(self, partition_rows):
        self._partitions = partition_rows
        self.sample_sqls, self.sample_params = [], []

    def query(self, sql, job_config=None):
        if "INFORMATION_SCHEMA.PARTITIONS" in sql:
            return list(self._partitions)
        self.sample_sqls.append(sql)
        self.sample_params.append([p.name for p in job_config.query_parameters])
        raise BadRequest("boom")


def _original_attempts(base, where, order, require):
    attempts = []
    if where:
        attempts.append(f"{base} {where}{order} LIMIT @n")
        if order:
            attempts.append(f"{base} {where} LIMIT @n")
    if not require:
        if order:
            attempts.append(f"{base}{order} LIMIT @n")
        attempts.append(f"{base} LIMIT @n")
    return attempts


@pytest.mark.parametrize("where_on", [False, True])
@pytest.mark.parametrize("order_on", [False, True])
@pytest.mark.parametrize("require", [False, True])
def test_fetch_last_rows_attempt_cascade(where_on, order_on, require, notes):
    client = CascadeClient([{"partition_id": "20260101"}])
    meta = {"partition_column": "pc" if where_on else "",
            "partition_column_type": "DATE" if where_on else "",
            "require_partition_filter": require,
            "temporal_columns": [("updt_ts", "TIMESTAMP")] if order_on else []}
    rows = gen.fetch_last_rows(client, "src-proj", "src_ds", "t", meta, n=7,
                               exclude=frozenset({"sec"}), notify=notes)
    assert rows == []
    base = "SELECT * EXCEPT (`sec`) FROM `src-proj.src_ds.t`"
    where = "WHERE `pc` = @part_day" if where_on else ""
    order = " ORDER BY `updt_ts` DESC" if order_on else ""
    if where_on:  # DATE partition also serves as the ORDER BY fallback
        order = order or " ORDER BY `pc` DESC"
    assert client.sample_sqls == _original_attempts(base, where, order, require), \
        (where_on, order_on, require, client.sample_sqls)
    for sql, params in zip(client.sample_sqls, client.sample_params):
        assert ("part_day" in params) == ("@part_day" in sql)
        assert "n" in params
    assert notes.warnings and "Could not sample rows" in notes.warnings[-1]


def test_latest_partition_filter_typed_params():
    client = CascadeClient([{"partition_id": "20260315"}])

    def meta(pct):
        return {"partition_column": "pc", "partition_column_type": pct,
                "require_partition_filter": False, "temporal_columns": []}

    clause, params = gen._latest_partition_filter(client, "src-proj", "src_ds",
                                                  "t", meta("DATE"))
    assert clause == "WHERE `pc` = @part_day"
    assert [(p.name, p.type_, str(p.value)) for p in params] == \
        [("part_day", "DATE", "2026-03-15")]

    clause, params = gen._latest_partition_filter(client, "src-proj", "src_ds",
                                                  "t", meta("TIMESTAMP"))
    assert clause == "WHERE `pc` >= @part_start AND `pc` < @part_end"
    assert [(p.name, p.type_) for p in params] == \
        [("part_start", "TIMESTAMP"), ("part_end", "TIMESTAMP")]
    assert params[0].value == datetime(2026, 3, 15, tzinfo=timezone.utc)
    assert params[1].value == datetime(2026, 3, 16, tzinfo=timezone.utc)

    clause, params = gen._latest_partition_filter(client, "src-proj", "src_ds",
                                                  "t", meta("DATETIME"))
    assert clause == "WHERE `pc` >= @part_start AND `pc` < @part_end"
    assert [(p.name, p.type_) for p in params] == \
        [("part_start", "DATETIME"), ("part_end", "DATETIME")]
    assert params[0].value == datetime(2026, 3, 15)  # naive for DATETIME columns

    assert gen._latest_partition_filter(client, "src-proj", "src_ds", "t",
                                        meta("INT64")) == ("", [])
    bad = CascadeClient([{"partition_id": "2026031507"}])  # hourly granularity
    assert gen._latest_partition_filter(bad, "src-proj", "src_ds", "t",
                                        meta("DATE")) == ("", [])
