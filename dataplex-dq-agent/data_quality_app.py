# Copyright 2026 Google LLC
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#     https://www.apache.org/licenses/LICENSE-2.0
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""Dataplex Auto-DQ (dataplex-dq) spec generator.

The LLM drafts an action plan and per-column rules from BigQuery profiling
results; scan ids, data_source blocks and file placement are deterministic.
Shared plumbing and the gated YAML validation/deploy path live in dq_common."""
import json
import os
import re
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timedelta, timezone
from typing import cast

import requests
import streamlit as st
from google.cloud import bigquery

from dq_common import (
    MAX_SCAN_ID_LEN, call_fuelix, collect_repo_scan_ids, fetch_table_metadata,
    finite_float, get_bq_client, load_abbreviations, parse_llm_json,
    pick_audit_column, read_target, render_file_header,
    render_preview_and_deploy, resolve_scan_id, select_tables, setup_page,
    show_target, sidebar_connection, sidebar_fuelix, sidebar_scan_settings,
    validate_scan_id, yaml_quote,
)

# --- Page + sidebar -----------------------------------------------------------
setup_page("Dataplex Auto-DQ Spec Generator")

environment, instance, source_project_id, repo_root = sidebar_connection()
project_id = st.sidebar.text_input("GCP Project ID (results table)",
                                   value="cio-datahub-governance-pr-8ba3")
dataset_id = st.sidebar.text_input(
    "Results Dataset ID", value="pr_datahub_01_ne1_edemm_dp_export_result",
    help="Dataset holding the scan RESULTS/profiling table.")
bq_location = st.sidebar.text_input("BigQuery Data Location", value="northamerica-northeast1")
profile_table_name = st.sidebar.text_input("Profiling Results Table",
                                           value="pr_datahub_01_ne1_governance_scan_results")
source_dataset_id = st.sidebar.text_input(
    "Source Dataset (scanned data)", value="ent_actvn",
    help="Dataset of the tables being scanned; drives the output file name "
         "and each scan's data_source.dataset_id.")

fuelix_api_key, model_name = sidebar_fuelix(
    "Model (FuelIX)", "Chat models exposed by the FuelIX gateway.")
gen_descriptions = st.sidebar.checkbox(
    "Table/column descriptions (review current, optionally generate via LLM)",
    value=False,
    help="Shows the descriptions currently stored on the BigQuery tables and asks "
         "whether to use them as-is (no LLM call), generate new ones, or fill only "
         "the gaps. Generating samples the last 1000 rows per table; per-column "
         "summaries and 20 rows go to the model (results saved to "
         "descriptions/*.xlsx) and are used alongside profiling stats for rule "
         "generation. When off, only Dataplex profiling statistics are used and no "
         "table data is sent to the LLM. Columns carrying a BigQuery policy tag "
         "are always excluded from every LLM payload.")
use_collibra, collibra_url, collibra_domain = False, "", ""
if gen_descriptions:
    use_collibra = st.sidebar.checkbox(
        "Enrich with Collibra glossary", value=False,
        help="Imports Business Terms (acronym, full name, definition) via the "
             "Collibra API — auth key read from Secret Manager — and injects the "
             "terms matched to table/column name segments into the description "
             "prompts.")
    if use_collibra:
        collibra_url = st.sidebar.text_input("Collibra URL",
                                             value="https://telus.collibra.com")
        collibra_domain = st.sidebar.text_input(
            "Collibra domain id (optional)", value="",
            help="Narrows the import to one glossary domain; blank = all "
                 "Business Terms org-wide.")
scan_cron, export_dataset, catalog_publishing_enabled = sidebar_scan_settings(
    "DQ Scan Settings", "0 7 * * *",
    "Header cron for a NEW dataset file; existing files keep theirs.")

governance_dir = os.path.join(repo_root, "edemm", environment, "governance")
job_prefix = instance
output_filename = f"{instance}_dqs_{source_dataset_id}.yaml"


# --- Profiling ------------------------------------------------------------------
@st.cache_data(ttl="15m", show_spinner=False)
def get_column_profiles(project: str, dataset: str, profile_table: str,
                        source_dataset: str, tables: tuple, location: str) -> list:
    """Latest profiling metrics per column for the target tables."""
    query = f"""
    WITH RankedProfiles AS (
      SELECT data_source.table_id AS table_name, column_name, column_type,
             column_mode, percent_null, percent_unique, min_value, max_value,
             average_value, standard_deviation, top_n,
             ROW_NUMBER() OVER (PARTITION BY data_source.table_id, column_name
                                ORDER BY job_start_time DESC) AS rank
      FROM `{project}.{dataset}.{profile_table}`
      WHERE data_source.dataset_id = @source_dataset_id
        AND data_source.table_id IN UNNEST(@table_names)
    )
    SELECT * FROM RankedProfiles WHERE rank = 1
    """
    job_config = bigquery.QueryJobConfig(query_parameters=[
        bigquery.ScalarQueryParameter("source_dataset_id", "STRING", source_dataset),
        bigquery.ArrayQueryParameter("table_names", "STRING", list(tables)),
    ])
    return [{**{k: row[k] for k in (
                "table_name", "column_name", "column_type", "column_mode",
                "percent_null", "percent_unique", "min_value", "max_value",
                "average_value", "standard_deviation")},
             "top_n": [{"value": i.get("value"), "count": i.get("count"),
                        "percent": i.get("percent")}
                       for i in (row["top_n"] if isinstance(row["top_n"], list) else [])]}
            for row in get_bq_client(project, location).query(query, job_config=job_config)]


# --- Policy-tag firewall: tagged columns never reach any LLM payload -------------
@st.cache_data(ttl="15m", show_spinner=False)
def fetch_bq_metadata(project: str, dataset: str, location: str, tables: tuple) -> dict:
    """{table: {"tagged": (dotted policy-tagged paths, ...), "table_description":
    str|None, "columns": {dotted path: description|None}}}. One threaded
    get_table per table (only the Tables API exposes policy tags/descriptions);
    raises on lookup failure so callers fail closed."""
    client = get_bq_client(project, location)

    def walk(fields, prefix=""):
        for f in fields:
            yield (prefix + f.name, f.description or None,
                   bool(f.policy_tags and f.policy_tags.names))
            yield from walk(f.fields or (), prefix + f.name + ".")

    def info(table: str) -> dict:
        t = client.get_table(f"{project}.{dataset}.{table}")
        entries = list(walk(t.schema))
        return {"tagged": tuple(path for path, _d, is_tagged in entries if is_tagged),
                "table_description": t.description or None,
                "columns": {path: d for path, d, _t in entries}}

    with ThreadPoolExecutor(max_workers=8) as pool:
        return dict(zip(tables, pool.map(info, tables)))


def _is_tagged(column: str, tagged_paths: tuple) -> bool:
    """True when the column is a tagged field or nested under/above one."""
    return any(column == t or column.startswith(t + ".") or t.startswith(column + ".")
               for t in tagged_paths)


def _label_list(labels: list) -> str:
    head = ", ".join(f"`{x}`" for x in labels[:60])
    return head + (f" … and {len(labels) - 60} more" if len(labels) > 60 else "")


def get_profiles_for_llm(tables: list) -> tuple[list, dict]:
    """(profiles, tagged): latest profiling rows minus policy-tagged columns —
    the single source of profiling data for LLM prompts — plus the tagged-path
    map for the description firewall and the reports."""
    profiles = get_column_profiles(project_id, dataset_id, profile_table_name,
                                   source_dataset_id, tuple(tables), bq_location)
    bq_meta = fetch_bq_metadata(source_project_id, source_dataset_id, bq_location,
                                tuple(tables))
    tagged = {t: info["tagged"] for t, info in bq_meta.items()}
    kept = [p for p in profiles
            if not _is_tagged(str(p.get("column_name")),
                              tagged.get(str(p.get("table_name")), ()))]
    excluded = [f"{t}.{c}" for t in sorted(tagged) for c in tagged[t]]
    if excluded:
        st.info("Policy-tagged columns excluded from every LLM payload (no sample "
                "data, no profiling statistics, no descriptions): "
                + _label_list(excluded))
    return kept, tagged


def existing_descriptions_from(bq_meta: dict) -> dict:
    """bq_meta reshaped to the descriptions shape used everywhere; policy-tagged
    columns are nulled so the firewall holds even when the current metadata is
    chosen for the next steps."""
    return {table: {
        "table_description": info["table_description"],
        "columns": {c: (None if _is_tagged(c, info["tagged"]) else d)
                    for c, d in info["columns"].items()}}
        for table, info in bq_meta.items()}


# --- Optional description generation (TELUS GenAI prompt standards) --------------
SAMPLE_ROW_COUNT = 1000     # window of most recent rows fetched per table
_EVIDENCE_ROWS = 20         # full rows from that window forwarded to the model
_EVIDENCE_TOP_VALUES = 12   # most frequent values per column forwarded
_EVIDENCE_CELL_CAP = 80     # cell truncation before prompting
UNABLE_PHRASE = "Unable to Generate Description"  # LLM sentinel; stored as null
_UNABLE_RE = re.compile(re.escape(UNABLE_PHRASE) + r"\.?", re.IGNORECASE)


def _null_if_unable(text) -> str | None:
    """Stored value for one LLM description: None for blank or the sentinel."""
    text = str(text or "").strip()
    return None if not text or _UNABLE_RE.fullmatch(text) else text


def _latest_partition_filter(project: str, dataset: str, table: str,
                             location: str, meta: dict) -> str:
    """WHERE clause pinning the newest partition via a zero-scan
    INFORMATION_SCHEMA.PARTITIONS lookup; '' unless day granularity resolves."""
    pc, pct = meta["partition_column"], meta["partition_column_type"]
    if not pc:
        return ""
    query = f"""
    SELECT partition_id
    FROM `{project}.{dataset}.INFORMATION_SCHEMA.PARTITIONS`
    WHERE table_name = @table AND partition_id IS NOT NULL
      AND partition_id NOT IN ('__NULL__', '__UNPARTITIONED__')
    ORDER BY partition_id DESC LIMIT 1
    """
    try:
        rows = list(get_bq_client(project, location).query(
            query, job_config=bigquery.QueryJobConfig(query_parameters=[
                bigquery.ScalarQueryParameter("table", "STRING", table)])))
    except Exception:
        return ""
    pid = str(rows[0]["partition_id"]) if rows else ""
    if len(pid) != 8 or not pid.isdigit():
        return ""
    day = f"{pid[:4]}-{pid[4:6]}-{pid[6:]}"
    if pct == "DATE":
        return f"WHERE `{pc}` = DATE '{day}'"
    if pct in ("TIMESTAMP", "DATETIME"):
        next_day = (date.fromisoformat(day) + timedelta(days=1)).isoformat()
        return f"WHERE `{pc}` >= '{day}' AND `{pc}` < '{next_day}'"
    return ""


def fetch_last_rows(project: str, dataset: str, table: str, location: str,
                    meta: dict, n: int = SAMPLE_ROW_COUNT,
                    exclude: frozenset = frozenset()) -> list:
    """Most recent n rows as dicts (best effort): newest partition, ordered by
    the audit/temporal column; retried without ORDER BY, then without the
    partition filter (unless the table requires one); [] when every attempt
    fails. `exclude` (policy-tagged top-level columns) is dropped in the SELECT
    itself, so protected data is never even fetched."""
    where = _latest_partition_filter(project, dataset, table, location, meta)
    order_col = pick_audit_column(meta["temporal_columns"])
    if not order_col and meta["partition_column_type"] in ("DATE", "TIMESTAMP", "DATETIME"):
        order_col = meta["partition_column"]
    order = f" ORDER BY `{order_col}` DESC" if order_col else ""
    select = ("* EXCEPT (" + ", ".join(f"`{c}`" for c in sorted(exclude)) + ")"
              if exclude else "*")
    base = f"SELECT {select} FROM `{project}.{dataset}.{table}`"
    attempts = dict.fromkeys(  # most-specific first, deduped when where/order are ''
        f"{base}{' ' + w if w else ''}{o} LIMIT {n}"
        for w, o in ((where, order), (where, ""), ("", order), ("", ""))
        if w or not meta["require_partition_filter"])
    client = get_bq_client(project, location)
    for sql in attempts:
        try:
            return [dict(row) for row in client.query(sql)]
        except Exception:
            continue
    st.warning(f"Could not sample rows from `{table}` — its descriptions will rely "
               "on profiling statistics and names only.")
    return []


def build_sample_evidence(rows: list, max_rows: int = _EVIDENCE_ROWS,
                          top_values: int = _EVIDENCE_TOP_VALUES,
                          cell_cap: int = _EVIDENCE_CELL_CAP) -> dict:
    """Compact prompt-safe digest: per-column non-null counts and most frequent
    values (over all rows), plus the newest max_rows full rows, cells truncated."""
    if not rows:
        return {}

    def clip(value):
        text = "" if value is None else str(value)
        return text[:cell_cap] + ("..." if len(text) > cell_cap else "")

    columns = {}
    for name in rows[0]:
        values = [row.get(name) for row in rows]
        top = Counter(clip(v) for v in values if v is not None).most_common(top_values)
        columns[name] = {"non_null": sum(v is not None for v in values),
                         "rows_sampled": len(values),
                         "top_values": [{"value": k, "count": c} for k, c in top]}
    return {"columns": columns,
            "recent_rows": [{k: clip(v) for k, v in row.items()} for row in rows[:max_rows]]}


# --- Optional Collibra glossary enrichment ----------------------------------------
COLLIBRA_SECRET_PROJECT = "cto-collibra-insights-pr-2267"
COLLIBRA_SECRET_NAME = "collibra_api_key"
_BUSINESS_TERM_TYPE_ID = "00000000-0000-0000-0000-000000011001"  # packaged Business Term
_GLOSSARY_MAX_LINES = 40
_CONTEXT_CHAR_CAP = 80_000  # keep description prompts inside gateway limits


@st.cache_data(ttl="1h", show_spinner=False)
def get_collibra_auth_key(project: str = COLLIBRA_SECRET_PROJECT,
                          secret: str = COLLIBRA_SECRET_NAME) -> str:
    """auth_key from Google Cloud Secret Manager (latest version)."""
    from google.cloud import secretmanager
    client = secretmanager.SecretManagerServiceClient()
    name = f"projects/{project}/secrets/{secret}/versions/latest"
    return client.access_secret_version({"name": name}).payload.data.decode("UTF-8")


def _collibra_auth_options(base_url: str, key: str) -> list:
    """Candidate requests-kwargs for the unknown key scheme, in probe order:
    Bearer, HTTP Basic (user:pass), pre-encoded Basic, session-login cookies."""
    options: list[dict] = [{"headers": {"Authorization": f"Bearer {key}"}},
                           {"headers": {"Authorization": f"Basic {key}"}}]
    if ":" in key:
        user, password = key.split(":", 1)
        options.insert(1, {"auth": (user, password)})
        try:
            resp = requests.post(f"{base_url.rstrip('/')}/rest/2.0/auth/sessions",
                                 json={"username": user, "password": password},
                                 timeout=30)
            if resp.ok and resp.cookies:
                options.append({"cookies": resp.cookies.get_dict()})
        except Exception:
            pass
    return options


@st.cache_data(ttl="1h", show_spinner=False)
def _pick_collibra_auth(base_url: str, auth_key: str) -> dict:
    """First auth scheme the API accepts, probed with a 1-row read."""
    statuses = []
    for kwargs in _collibra_auth_options(base_url, auth_key):
        resp = requests.get(f"{base_url.rstrip('/')}/rest/2.0/assets",
                            params={"limit": 1}, timeout=30, **kwargs)
        if resp.ok:
            return kwargs
        statuses.append(str(resp.status_code))
    raise RuntimeError("Collibra rejected every auth scheme "
                       f"(Bearer/Basic/session): HTTP {', '.join(statuses)}")


@st.cache_data(ttl="1h", show_spinner=False)
def fetch_collibra_glossary(base_url: str, domain_id: str, auth_key: str) -> dict:
    """{term_lower: {id, name, full}} for every Business Term, paged 1000/call
    until the reported total (optionally narrowed to one domain). Raises on
    HTTP/auth failure — the caller degrades gracefully."""
    auth = _pick_collibra_auth(base_url, auth_key)
    terms, offset, total = {}, 0, 1
    params = {"typeIds": _BUSINESS_TERM_TYPE_ID, "limit": 1000}
    if domain_id:
        params["domainId"] = domain_id
    while offset < total:
        resp = requests.get(f"{base_url.rstrip('/')}/rest/2.0/assets",
                            params={**params, "offset": offset}, timeout=60, **auth)
        resp.raise_for_status()
        page = resp.json()
        total = int(page.get("total") or 0)
        results = page.get("results") or []
        if not results:
            break
        for asset in results:
            name = str(asset.get("name") or "")
            if name:
                terms[name.lower()] = {"id": str(asset.get("id") or ""), "name": name,
                                       "full": str(asset.get("displayName") or name)}
        offset += len(results)
    return terms


@st.cache_data(ttl="1h", show_spinner=False)
def fetch_collibra_definitions(base_url: str, auth_key: str, asset_ids: tuple) -> dict:
    """{asset_id: definition text} for matched terms (capped 60); HTML stripped,
    per-asset failures skipped."""
    auth = _pick_collibra_auth(base_url, auth_key)
    defs = {}
    for asset_id in asset_ids[:60]:
        try:
            resp = requests.get(f"{base_url.rstrip('/')}/rest/2.0/attributes",
                                params={"assetId": asset_id, "limit": 20},
                                timeout=30, **auth)
            resp.raise_for_status()
            for attr in resp.json().get("results") or []:
                type_name = str((attr.get("type") or {}).get("name") or "").lower()
                if "definition" in type_name or "description" in type_name:
                    text = re.sub(r"<[^>]+>", " ", str(attr.get("value") or ""))
                    defs[asset_id] = re.sub(r"\s+", " ", text).strip()
                    break
        except Exception:
            continue
    return defs


def _name_tokens(names) -> set:
    """Alphanumeric segments (len >= 2) of the given table/column names."""
    return {t for n in names for t in re.split(r"[^a-zA-Z0-9]+", n.lower()) if len(t) >= 2}


def collibra_hint(glossary: dict, base_url: str, auth_key: str, *names: str) -> str:
    """Up to _GLOSSARY_MAX_LINES 'TERM = Full Name — Definition' lines for
    glossary terms matching the given name segments; '' if none."""
    matched = [glossary[t] for t in sorted(_name_tokens(names))
               if t in glossary][:_GLOSSARY_MAX_LINES]
    if not matched:
        return ""
    defs = fetch_collibra_definitions(base_url, auth_key,
                                      tuple(m["id"] for m in matched if m["id"]))
    return "\n".join(f"{m['name']} = {m['full']}"
                     + (f" — {defs[m['id']]}" if defs.get(m["id"]) else "")
                     for m in matched)


def _reverse_abbreviations() -> dict:
    """{abbreviation: 'full word(s)'} from abbreviations.csv; drop-tokens
    (empty abbreviation) skipped, shared abbreviations joined with ' / '."""
    reverse: dict = {}
    for full, abbr in load_abbreviations().items():
        if abbr:
            reverse[abbr] = f"{reverse[abbr]} / {full}" if abbr in reverse else full
    return reverse


def abbrev_hint(*names: str) -> str:
    """Up to _GLOSSARY_MAX_LINES 'abbr = full word(s)' lines for abbreviations
    matching the given name segments; '' if none."""
    reverse = _reverse_abbreviations()
    matched = [f"{t} = {reverse[t]}" for t in sorted(_name_tokens(names)) if t in reverse]
    return "\n".join(matched[:_GLOSSARY_MAX_LINES])


# Constraints below follow the TELUS GenAI description prompt standards
# (table-level: 16 constraints; column-level: 15 constraints, batch-adapted).
_SYS_TABLE_DESC = """You are a TELUS data steward generating a business description for a BigQuery table.
You will receive the table name, Dataplex column profiling statistics, and evidence sampled from the table's most recent rows. Use them only to infer business meaning.
If you are unable to generate a description, respond with exactly: Unable to Generate Description

Here are examples of good table descriptions to use as a reference:
Example 1: This table captures the different types of equipment classification. A product equipment classification is a secondary level grouping of the equipment, organized by its use.
Example 2: This table holds outgoing notifications to the TELUS customers for different marketing campaigns.

Your response MUST satisfy ALL the following constraints:
1. The response MUST NOT repeat the provided table name within the description.
2. The response MUST start with the following words: This table contains.
3. The response MUST be clear, concise, and complete.
4. The response MUST explain any key business terms as part of the description.
5. The response MUST be stated in the present tense, in the form of a complete descriptive sentence.
6. The response MUST NOT include jargon or abbreviations.
7. If the response references an acronym, the acronym MUST be spelled out.
8. The response MUST be provided from a business perspective, independent from any database or software context.
9. The response MUST NOT include any potential example/sample values.
10. The response MUST NOT include any details about potential attributes or columns collected or stored in the table.
11. The response MUST EXCLUDE any pleasantries.
12. The response MUST EXCLUDE any external references or citations.
13. The response MUST be confident. DO NOT use words such as "this table may contain" or "this table is likely to contain".
14. The response MUST NOT provide an opinion as to why this information is valuable.
15. The response MUST NOT include filler words and phrases, such as "in the organization" or any synonyms. If this information is required for a complete definition, use TELUS rather than a generic form such as organization or company.
16. If the response references the company name TELUS, ensure that TELUS is capitalized.

When generating the description, also consider:
- Some tables start with the prefix bq_; ignore the prefix and use the rest of the name (bq_customer_info -> customer_info).
- Some tables end with the suffix _dim; ignore the suffix (team_member_dim -> team_member).
- Table names are never about an individual person, so avoid references to specific individuals.
- The profiling statistics and sampled values are context only and MUST NOT appear in the output.
Return ONLY the description text (or the exact fallback phrase) — no preamble, no markdown."""

_SYS_COL_DESC = """You are a TELUS data steward generating business descriptions for the columns of a BigQuery table.
You will receive the table name, its generated table description, the list of columns to describe, Dataplex column profiling statistics, and evidence sampled from the table's most recent rows. Use them only to infer business meaning.
Return ONLY one JSON object (no markdown, no prose) of this shape, covering EVERY listed column:
{"columns": {"<column_name>": "<description>"}}
Use the exact phrase "Unable to Generate Description" as the value for any column you cannot describe.

Here are examples of good column descriptions to use as a reference:
Example 1: This column contains the cardinal direction from which the wind blows in an abbreviated form. Wind directions are always expressed as from whence the wind blows meaning that a North wind blows from North to South.
Example 2: The date printed on the Bill Document indicating the extraction date of Customer data from the billing system.
Example 3: A yes/no indicator identifying whether or not the party is hearing impaired. This field only applies to Parties that are Individuals.
Example 4: The extended customer name such as trade name for the sole proprietorship business customer. It is the trade/firm name that followed the operating type such as Operating As (O/A), Doing Business as (D/B).
Example 5: A code used to represent the discount level that the customer is entitled to for various products and services. This only applies to those customers that are also team members.
Example 6: This column contains an ISO defined unique identifier of a currency. In this instance, it represents the currency in which the customers' invoices are presented.

Every description MUST satisfy ALL the following constraints:
1. The description MUST NOT repeat the provided table name or column name.
2. The description MUST start with the following words: This column contains.
3. The description MUST be clear, concise, and complete.
4. The description MUST explain any key business terms.
5. The description MUST be stated in the present tense, in the form of a complete descriptive sentence.
6. The description MUST NOT include jargon or abbreviations.
7. If the description references an acronym, the acronym MUST be spelled out.
8. The description MUST be provided from a business perspective, independent from any database or software context.
9. The description MUST NOT include any potential example/sample values.
10. The description MUST EXCLUDE any pleasantries.
11. The description MUST EXCLUDE any external references or citations.
12. The description MUST be confident. DO NOT use words such as "this column may contain" or "this column is likely to contain".
13. The description MUST NOT provide an opinion as to why this information is valuable.
14. The description MUST NOT include filler words and phrases, such as "in the organization" or any synonyms. If this information is required for a complete definition, use TELUS rather than a generic form such as organization or company.
15. If the description references the company name TELUS, ensure that TELUS is capitalized.

When generating the column descriptions, also consider:
- Columns ending with the suffix _ind are usually indicators (e.g. INDVDL_HEARING_IMPAIRED_IND indicates whether or not the individual is hearing impaired).
- The profiling statistics and sampled values are context only and MUST NOT appear in the output."""


def _table_profiles(profiles: list, table: str) -> list:
    return [{k: v for k, v in p.items() if k != "table_name"}
            for p in profiles if p.get("table_name") == table]


def _jsonc(obj) -> str:  # compact prompt JSON; default=str for dates/decimals
    return json.dumps(obj, separators=(",", ":"), default=str)


def generate_descriptions(tables: list, profiles: list, tagged: dict,
                          existing: dict | None = None,
                          fill_only_missing: bool = False) -> dict:
    """{table: {"table_description": str|None, "columns": {col: str|None}}}:
    two FuelIX calls per table grounded in the last SAMPLE_ROW_COUNT rows and
    the policy-filtered profiling stats. `existing` (already policy-nulled)
    rides along as reference context; fill_only_missing keeps it verbatim and
    generates only the gaps (fully described tables skip sampling and the LLM).
    Tagged columns, blank/sentinel replies and per-table failures store null."""
    existing = existing or {}
    meta = fetch_table_metadata(source_project_id, source_dataset_id,
                                bq_location, tuple(tables))
    glossary, collibra_key = {}, ""
    if use_collibra:
        try:
            collibra_key = get_collibra_auth_key()
            glossary = fetch_collibra_glossary(collibra_url, collibra_domain, collibra_key)
            st.caption(f"Collibra glossary imported: {len(glossary)} business terms.")
        except Exception as e:
            st.warning(f"Collibra glossary unavailable ({e}) — continuing without it.")
    out = {}
    for table in tables:
        tagged_paths = tagged.get(table, ())
        cur = existing.get(table) or {"table_description": None, "columns": {}}
        cur_cols = {c: d for c, d in cur["columns"].items() if d}
        if fill_only_missing and cur["table_description"] and not any(
                not d and not _is_tagged(c, tagged_paths)
                for c, d in cur["columns"].items()):
            out[table] = {"table_description": cur["table_description"],
                          "columns": dict(cur["columns"])}  # fully described: no LLM
            continue
        rows = fetch_last_rows(source_project_id, source_dataset_id, table,
                               bq_location, meta[table],
                               exclude=frozenset(t.split(".")[0] for t in tagged_paths))
        stats = _table_profiles(profiles, table)
        stats_json = _jsonc(stats)
        evidence = build_sample_evidence(rows)
        col_names = (list(evidence["columns"]) if evidence
                     else sorted({str(s["column_name"]) for s in stats}))

        def render_context(ev):
            return (f"TABLE NAME: {table}\nPROFILING STATISTICS:\n{stats_json}\n"
                    f"SAMPLE EVIDENCE (from the last {SAMPLE_ROW_COUNT} rows):\n"
                    + (_jsonc(ev) if ev else "(no rows sampled)"))

        # Wide tables can blow the gateway request cap; shrink evidence stepwise.
        context = render_context(evidence)
        if len(context) > _CONTEXT_CHAR_CAP:
            evidence = build_sample_evidence(rows, max_rows=5, top_values=5, cell_cap=40)
            context = render_context(evidence)
        if len(context) > _CONTEXT_CHAR_CAP and evidence:
            evidence.pop("recent_rows", None)
            context = render_context(evidence)
        if len(context) > _CONTEXT_CHAR_CAP:
            context = render_context({})[:_CONTEXT_CHAR_CAP]  # stats-only, hard-capped
        hint = abbrev_hint(table, *col_names)
        if hint:
            context += ("\nTELUS ABBREVIATION GLOSSARY (abbreviations.csv; matched "
                        "to this table's name/column segments):\n" + hint)
        ch = collibra_hint(glossary, collibra_url, collibra_key,
                           table, *col_names) if glossary else ""
        if ch:
            context += ("\nTELUS BUSINESS GLOSSARY (Collibra; matched to this "
                        "table's name/column segments — reference hints for "
                        "interpreting abbreviations, not data):\n" + ch)
        ref = {k: v for k, v in (("table_description", cur["table_description"]),
                                 ("columns", cur_cols)) if v}
        if ref:
            context += ("\nEXISTING BIGQUERY DESCRIPTIONS (current metadata — "
                        "reference only, may be incomplete or outdated):\n"
                        + _jsonc(ref)[:20_000])
        cols_to_ask = ([c for c in col_names if not cur_cols.get(c)]
                       if fill_only_missing else col_names)
        try:
            table_desc = (cur["table_description"]
                          if fill_only_missing and cur["table_description"]
                          else _null_if_unable(call_fuelix(
                              _SYS_TABLE_DESC,
                              "Generate the table description.\n" + context,
                              fuelix_api_key, model_name)))
            col_map: dict[str, str | None] = {}
            if cols_to_ask:
                col_prompt = (f"COLUMNS TO DESCRIBE: {json.dumps(cols_to_ask)}\n"
                              f"TABLE DESCRIPTION: {table_desc or ''}\n" + context +
                              "\nGenerate the JSON object of column descriptions.")
                data = parse_llm_json(call_fuelix(_SYS_COL_DESC, col_prompt,
                                                  fuelix_api_key, model_name))
                found = data.get("columns", data) if isinstance(data, dict) else {}
                col_map = {c: _null_if_unable(found.get(c)) for c in cols_to_ask}
        except Exception as e:
            st.warning(f"Descriptions for `{table}` failed: {e}")
            table_desc = cur["table_description"] if fill_only_missing else None
            col_map = dict.fromkeys(cols_to_ask)
        if fill_only_missing:  # keep the current descriptions verbatim
            for c, d in cur_cols.items():
                col_map.setdefault(c, d)
        for path in tagged_paths:  # visible in Excel/report, never sent to the LLM
            col_map.setdefault(path, None)
        out[table] = {"table_description": table_desc, "columns": col_map}
    return out


_XLSX_HEADER = ["dataset", "table_name", "column_name", "description", "model",
                "generated_ts"]


def descriptions_xlsx_path() -> str:
    return os.path.join(os.path.dirname(__file__), "descriptions",
                        f"{instance}_dq_descriptions_{source_dataset_id}.xlsx")


def save_descriptions_xlsx(path: str, descriptions: dict) -> bool:
    """Upsert into the repo Excel keyed on (dataset, table_name, column_name):
    one table-level row (empty column_name) plus one row per column. Rows for
    other datasets/tables are preserved and duplicate key rows all kept in sync.
    Targets the 'descriptions' sheet by name, validates its header, writes
    atomically; returns False (with st.error) instead of raising on I/O problems
    such as the workbook being open in Excel."""
    try:
        import openpyxl
        from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
    except ImportError:
        st.error("openpyxl is not installed — run `pip install openpyxl` to save "
                 "descriptions to Excel.")
        return False

    def clean(value):
        # None passes through so a null description stays a truly blank cell
        # (and the upsert clears any previously saved placeholder text).
        return None if value is None else ILLEGAL_CHARACTERS_RE.sub("", str(value))

    def cell_key(value) -> str:
        return "" if value is None else str(value).strip()

    ts = datetime.now(timezone.utc).isoformat(timespec="seconds")
    fresh = {(cell_key(source_dataset_id), cell_key(table), cell_key(col)): clean(desc)
             for table, info in descriptions.items()
             for col, desc in [("", info["table_description"]), *info["columns"].items()]}
    tmp = path + ".tmp"
    try:
        exists = os.path.exists(path)
        wb = openpyxl.load_workbook(path) if exists else openpyxl.Workbook()
        if not exists and wb.active is not None:
            wb.remove(wb.active)
        if "descriptions" in wb.sheetnames:
            ws = wb["descriptions"]
            header = [c.value for c in next(ws.iter_rows(
                min_row=1, max_row=1, max_col=len(_XLSX_HEADER)))]
            if all(v is None for v in header):
                for i, name in enumerate(_XLSX_HEADER, start=1):
                    ws.cell(row=1, column=i, value=name)
            elif header != _XLSX_HEADER:
                st.error(f"Sheet 'descriptions' in `{path}` has an unexpected "
                         f"header {header} — expected {_XLSX_HEADER}. Fix or "
                         "delete the file, then regenerate.")
                return False
        else:
            ws = wb.create_sheet("descriptions")
            ws.append(_XLSX_HEADER)
        updated = set()
        for row in ws.iter_rows(min_row=2, max_col=len(_XLSX_HEADER)):
            key = (cell_key(row[0].value), cell_key(row[1].value), cell_key(row[2].value))
            if key in fresh:
                # assign via .value: ws.cell(value=None) skips the write and
                # would leave stale text behind instead of blanking the cell.
                cells = cast("tuple[Cell, ...]", row)
                cells[3].value, cells[4].value, cells[5].value = fresh[key], model_name, ts
                updated.add(key)
        for (ds, table, col), desc in fresh.items():
            if (ds, table, col) not in updated:
                ws.append([ds, table, col, desc, model_name, ts])
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        wb.save(tmp)
        os.replace(tmp, path)
        return True
    except Exception as e:
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except OSError:
            pass
        st.error(f"Could not write `{path}` ({e}) — if the file is open in "
                 "Excel, close it and regenerate. The descriptions are still "
                 "available in this session.")
        return False


def flatten_descriptions(descriptions: dict) -> list:
    """Preview rows for st.dataframe (table-level row first per table)."""
    return [{"table": table, "column": col, "description": desc}
            for table, info in descriptions.items()
            for col, desc in [("(table)", info["table_description"]),
                              *info["columns"].items()]]


def null_description_labels(descriptions: dict, tagged: dict) -> tuple[list, list]:
    """(failed, policy_excluded) 'table' / 'table.column' labels for the null
    descriptions, split by cause for the post-generation report."""
    failed, excluded = [], []
    for table, info in descriptions.items():
        if info.get("table_description") is None:
            failed.append(table)
        for col, desc in info.get("columns", {}).items():
            if desc is None:
                (excluded if _is_tagged(col, tagged.get(table, ()))
                 else failed).append(f"{table}.{col}")
    return failed, excluded


def active_descriptions():
    """Descriptions for prompt injection — None while the feature is off, so a
    stale session/Excel can never leak into prompts; null descriptions pruned,
    so neither nulls nor policy-tagged names reach the action-plan/rule prompts."""
    if not gen_descriptions:
        return None
    pruned = {}
    for table, info in (st.session_state.get("descriptions") or {}).items():
        entry = {k: v for k, v in (
            ("table_description", info.get("table_description")),
            ("columns", {c: d for c, d in info.get("columns", {}).items() if d})) if v}
        if entry:
            pruned[table] = entry
    return pruned or None


# --- YAML rendering (dataplex-dq scan format) -------------------------------------
_I_DASH, _I_KEY, _I_SUB = " " * 14, " " * 16, " " * 18
_EXPECTATIONS = {
    "non_null_expectation", "uniqueness_expectation", "range_expectation",
    "set_expectation", "regex_expectation", "table_condition_expectation",
}


def _num(value) -> str:
    """Bare number when finite (whole floats without decimal), else a quoted
    scalar — so string bounds like '2020-01-01' can't corrupt the document."""
    n = finite_float(value)
    if n is None:
        return yaml_quote(value)
    return str(int(n)) if n == int(n) else str(n)


def rule_is_valid(rule) -> bool:
    """Renderable = name + dimension + a KNOWN expectation type; unknown types
    would otherwise render as an expectation-less rule Dataplex rejects."""
    return (isinstance(rule, dict) and bool(rule.get("name"))
            and bool(rule.get("dimension"))
            and (rule.get("expectation") or {}).get("type") in _EXPECTATIONS)


def _sanitize_rule(rule):
    """Coerce one LLM rule toward validity: Dataplex name charset, float
    threshold (dropped if not coercible), stripped column/dimension, list-typed
    set values."""
    if not isinstance(rule, dict):
        return rule
    if rule.get("name"):
        rule["name"] = re.sub(r"[^a-z0-9-]+", "-", str(rule["name"]).lower()).strip("-")
    if rule.get("column") is not None:
        rule["column"] = str(rule["column"]).strip()
    if rule.get("dimension") is not None:
        rule["dimension"] = str(rule["dimension"]).strip().upper()
    if rule.get("threshold") is not None:
        rule["threshold"] = finite_float(str(rule["threshold"]).strip().rstrip("%"))
        if rule["threshold"] is None:
            del rule["threshold"]
    exp = rule.get("expectation")
    if isinstance(exp, dict) and exp.get("type") == "set_expectation":
        values = exp.get("values")
        exp["values"] = values if isinstance(values, list) else ([] if values is None else [values])
    return rule


def render_rule(rule: dict) -> list:
    exp = rule.get("expectation") or {}
    etype, column = exp.get("type"), rule.get("column")
    lines = ([_I_DASH + f"- column: {yaml_quote(column)}",
              _I_KEY + f"name: {yaml_quote(rule['name'])}"]
             if column else [_I_DASH + f"- name: {yaml_quote(rule['name'])}"])
    lines.append(_I_KEY + f"dimension: {yaml_quote(rule['dimension'])}")
    threshold = finite_float(rule.get("threshold"))
    if etype != "table_condition_expectation" and threshold is not None:
        lines.append(_I_KEY + f"threshold: {threshold}")
    if etype in ("non_null_expectation", "uniqueness_expectation"):
        lines.append(_I_KEY + f"{etype}: true")
    elif etype == "range_expectation":
        lines.append(_I_KEY + "range_expectation:")
        for key in ("min_value", "max_value"):
            if exp.get(key) is not None:
                lines.append(_I_SUB + f"{key}: {_num(exp[key])}")
    elif etype == "set_expectation":
        lines.append(_I_KEY + "set_expectation:")
        lines.append(_I_SUB + "values: ["
                     + ", ".join(yaml_quote(v) for v in exp.get("values") or []) + "]")
        if exp.get("ignore_null"):
            lines.append(_I_SUB + "ignore_null: true")
    elif etype == "regex_expectation":
        lines.append(_I_KEY + "regex_expectation:")
        lines.append(_I_SUB + f"regex: {yaml_quote(exp.get('regex', ''))}")
        if exp.get("ignore_null"):
            lines.append(_I_SUB + "ignore_null: true")
    else:  # table_condition_expectation
        lines.append(_I_KEY + "table_condition_expectation:")
        lines.append(_I_SUB + f"sql_expression: {yaml_quote(exp.get('sql_expression', ''))}")
    return lines


def render_scan_block(scan_id, project_id_, dataset_id_, table_id_, rules) -> str:
    lines = [
        f"        {scan_id}:",
        "          data_source:",
        f"            project_id: {yaml_quote(project_id_)}",
        f"            dataset_id: {yaml_quote(dataset_id_)}",
        f"            table_id: {yaml_quote(table_id_)}",
        "          data_quality_spec:",
        "            rules:",
    ]
    for rule in rules:
        lines.extend(render_rule(rule))
    return "\n".join(lines)


# --- LLM steps ---------------------------------------------------------------------
_SYS_PLAN = """You are an expert Google Cloud Dataplex data quality engineer.
From the provided BigQuery column profiling data, draft a step-by-step Action Plan for a Dataplex DQ scan:
1. Identify columns suited to non_null_expectation (COMPLETENESS), uniqueness_expectation (UNIQUENESS), range/set/regex_expectation (VALIDITY), or table-level table_condition_expectation (VOLUME/FRESHNESS).
2. Justify each suggestion from the statistics (e.g. "range_expectation for antenna_face: values span 1-3 with no outliers"; "uniqueness_expectation for acct_id: percent_unique is 100%").
3. Do NOT write YAML/JSON yet — analysis and justification only, in clear markdown."""

_SYS_RULES = """You are an expert Google Cloud Dataplex data quality engineer.
From column profiling data, an action plan and user feedback, return ONLY one JSON object (no markdown, no prose) of this shape:
{"tables": {"<table_name>": {"rules": [{
  "column": "<column_name>",   // omit for table-level rules
  "name": "<rule-name>",       // lowercase letters, digits and hyphens ONLY
  "dimension": "COMPLETENESS|UNIQUENESS|VALIDITY|VOLUME|FRESHNESS",
  "threshold": 0.99,           // float; omit for table-level rules
  "expectation": {...}}]}}}    // exactly one of:
- {"type":"non_null_expectation"}                                        COMPLETENESS, threshold 0.99
- {"type":"uniqueness_expectation"}                                      UNIQUENESS, threshold 1.0
- {"type":"range_expectation","min_value":0,"max_value":9}               VALIDITY, threshold 0.99 (min and/or max)
- {"type":"set_expectation","values":["A","B"],"ignore_null":true}       VALIDITY, threshold 0.99
- {"type":"regex_expectation","regex":"^.{9}$","ignore_null":true}       VALIDITY, threshold 0.99
- {"type":"table_condition_expectation","sql_expression":"COUNT(*) > 0"} VOLUME/FRESHNESS; NO column, NO threshold
Use table_name keys exactly as given. Only reference columns present in the profiling data. Incorporate all user feedback."""


_DESC_LABEL = ("Table & Column Business Descriptions (use them to infer "
               "business rules the statistics alone cannot show):\n")


def generate_action_plan(profile_json, descriptions=None):
    prompt = ("Analyze this column profiling data and draft a data quality action plan:\n"
              + _jsonc(profile_json))
    if descriptions:
        prompt += "\n\n" + _DESC_LABEL + _jsonc(descriptions)
    return call_fuelix(_SYS_PLAN, prompt, fuelix_api_key, model_name)


def _read_upload(uploaded_file) -> str:
    """Best-effort text of an uploaded reference file; '' if binary."""
    if uploaded_file is None:
        return ""
    raw = uploaded_file.read()
    for enc in ("utf-8", "latin-1"):
        try:
            text = raw.decode(enc)
            if text.strip():
                return text
        except (UnicodeDecodeError, AttributeError):
            pass
    st.warning(f"'{getattr(uploaded_file, 'name', 'file')}' is not readable as text "
               "(binary formats are unsupported) — paste key requirements into the feedback box.")
    return ""


def generate_rules(profile_json, action_plan, hitl_feedback, uploaded_file=None,
                   descriptions=None):
    """{table_name: [rule_dict, ...]} from the model; sanitized per rule."""
    desc_section = _DESC_LABEL + _jsonc(descriptions) + "\n\n" if descriptions else ""
    prompt = (f"Profiling Data:\n{_jsonc(profile_json)}\n\n"
              + desc_section +
              f"Proposed Action Plan:\n{action_plan}\n\n"
              f"User Adjustments/HITL Feedback:\n{hitl_feedback}\n\n"
              f"Reference document contents (if any):\n{_read_upload(uploaded_file) or '(none provided)'}\n\n"
              "Return the JSON object of rules per table.")
    raw_text = call_fuelix(_SYS_RULES, prompt, fuelix_api_key, model_name)
    data = parse_llm_json(raw_text)
    if data is None:
        raise ValueError("Model did not return JSON. Raw response:\n" + raw_text)
    tables = data.get("tables", data) if isinstance(data, dict) else {}
    return {tbl: [_sanitize_rule(r) for r in
                  (spec.get("rules", []) if isinstance(spec, dict)
                   else spec if isinstance(spec, list) else [])]
            for tbl, spec in tables.items()}


def build_scan_plan(table_order, rules_by_table):
    """(plan, blocks, scan_ids): report rows plus rendered blocks and their ids
    for tables that passed scan-id validation and have valid rules."""
    forbidden = collect_repo_scan_ids(governance_dir, job_prefix)
    plan, blocks, ids = [], [], []
    for table in table_order:
        rules = [r for r in (rules_by_table.get(table) or []) if rule_is_valid(r)]
        scan_id = resolve_scan_id(job_prefix, source_dataset_id, table,
                                  forbidden | set(ids))
        ok, reason = validate_scan_id(scan_id, job_prefix, forbidden | set(ids))
        if ok and not rules:
            ok, reason = False, "no valid rules generated"
        if ok:
            blocks.append(render_scan_block(scan_id, source_project_id,
                                            source_dataset_id, table, rules))
            ids.append(scan_id)
        plan.append({"table": table, "scan_id": scan_id, "rules": len(rules),
                     "status": "ok" if ok else f"skip: {reason}"})
    return plan, blocks, ids


# --- UI flow -------------------------------------------------------------------------
st.write("### Target Selection")
st.caption(f"Env **{environment}** / **{instance}**, source project `{source_project_id}`, "
           f"output `edemm/{environment}/governance/{output_filename}`.")
table_names = select_tables(source_project_id, source_dataset_id, bq_location, "fetched_tables")

if table_names and gen_descriptions:
    st.write("---")
    st.write("### Optional Step: Table & Column Descriptions")
    st.caption("Fetch the descriptions currently stored on the BigQuery tables, "
               "review them, then choose: keep them for the next steps, generate "
               f"new ones with the LLM (per-column summaries and up to {_EVIDENCE_ROWS} "
               f"of the last {SAMPLE_ROW_COUNT} rows per table go to the model, held "
               "to the TELUS description standards), or a hybrid that only fills the "
               "gaps. The chosen set is saved to "
               f"`descriptions/{instance}_dq_descriptions_{source_dataset_id}.xlsx` "
               "and fed into the action plan and rule generation. Policy-tagged "
               "columns are excluded from every LLM payload, and descriptions the "
               "model cannot produce are stored as null.")
    if st.button("Fetch Current Descriptions"):
        with st.spinner("Reading current table metadata from BigQuery..."):
            try:
                bq_meta = fetch_bq_metadata(source_project_id, source_dataset_id,
                                            bq_location, tuple(table_names))
                st.session_state["existing_descriptions"] = existing_descriptions_from(bq_meta)
                st.session_state["desc_policy_tags"] = {t: i["tagged"]
                                                        for t, i in bq_meta.items()}
            except Exception as e:
                st.error(f"Error fetching current descriptions: {e}")
    if "existing_descriptions" in st.session_state:
        existing = st.session_state["existing_descriptions"]
        st.write("#### Current BigQuery descriptions")
        st.dataframe(flatten_descriptions(existing), width="stretch")
        col_descs = [d for i in existing.values() for d in i["columns"].values()]
        st.caption(f"{sum(1 for i in existing.values() if i['table_description'])} of "
                   f"{len(existing)} tables and {sum(1 for d in col_descs if d)} of "
                   f"{len(col_descs)} columns already have a description.")
        choice = st.radio(
            "How should descriptions for the next steps be produced?",
            ["Generate new descriptions with the LLM",
             "Use the current BigQuery descriptions (no LLM call)",
             "Hybrid: keep current descriptions, generate only the missing ones"])
        use_current = choice.startswith("Use")
        hybrid = choice.startswith("Hybrid")
        action_label = ("Use Current Descriptions" if use_current
                        else "Generate Missing Descriptions" if hybrid
                        else "Generate Descriptions")
        if st.button(action_label):
            try:
                if use_current:
                    st.session_state["descriptions"] = existing
                else:
                    with st.spinner("Sampling recent rows and drafting descriptions..."):
                        profiles, tagged = get_profiles_for_llm(table_names)
                        st.session_state["descriptions"] = generate_descriptions(
                            table_names, profiles, tagged,
                            existing=existing, fill_only_missing=hybrid)
                        st.session_state["desc_policy_tags"] = tagged
                st.session_state.pop("action_plan", None)
                st.session_state.pop("rules_by_table", None)
                xlsx_path = descriptions_xlsx_path()
                if save_descriptions_xlsx(xlsx_path, st.session_state["descriptions"]):
                    st.success(f"Descriptions saved to `{xlsx_path}`.")
            except Exception as e:
                st.error(f"Error preparing descriptions: {e}")
    if "descriptions" in st.session_state:
        st.write("#### Descriptions for the next steps")
        st.dataframe(flatten_descriptions(st.session_state["descriptions"]), width="stretch")
        failed, excluded = null_description_labels(
            st.session_state["descriptions"], st.session_state.get("desc_policy_tags", {}))
        if failed:
            st.warning("No description could be generated (stored as null) for: "
                       + _label_list(failed))
        if excluded:
            st.info("Excluded from LLM payloads by policy tags (description left "
                    "null): " + _label_list(excluded))

if table_names:
    st.write("---")
    st.write("### Step 1: Formulate Action Plan")
    if st.button("Generate Action Plan"):
        with st.spinner("Retrieving profiling statistics and drafting plan..."):
            try:
                profiles = get_profiles_for_llm(table_names)[0]
                st.session_state["profiles"] = profiles
                st.session_state["table_order"] = table_names
                st.session_state.pop("rules_by_table", None)
                if not profiles:
                    st.warning("No profiling data found for the selected table(s).")
                else:
                    st.session_state["action_plan"] = generate_action_plan(
                        profiles, active_descriptions())
            except Exception as e:
                st.error(f"Error generating action plan: {e}")

if "action_plan" in st.session_state:
    st.markdown("#### Proposed Action Plan")
    st.markdown(st.session_state["action_plan"])
    st.write("---")
    st.write("### Step 2: Feedback & Rule Generation")
    hitl_feedback = st.text_area(
        "Apply Custom Business Knowledge / Override Rules",
        value="The plan looks good. Please proceed.",
        help="e.g. 'Remove volume check', 'Add Antarctica to geo_country allowed set'.")
    uploaded_file = st.file_uploader("Upload Rules Reference (PDF, CSV, Excel, TXT)",
                                     type=["pdf", "csv", "xlsx", "xls", "txt"])
    if st.button("Generate DQ Rules"):
        with st.spinner("Incorporating feedback and generating rules..."):
            try:
                st.session_state["rules_by_table"] = generate_rules(
                    st.session_state["profiles"], st.session_state["action_plan"],
                    hitl_feedback, uploaded_file, active_descriptions())
            except Exception as e:
                st.error(f"Error generating rules: {e}")

if "rules_by_table" in st.session_state:
    plan, blocks, ids = build_scan_plan(
        st.session_state.get("table_order", table_names), st.session_state["rules_by_table"])
    st.write("---")
    st.write("### Step 3: Validate & Deploy")
    st.caption(f"Scan-id rules: length <= {MAX_SCAN_ID_LEN}, `^{job_prefix}_[a-z0-9_]+$`, "
               "no trailing separator, unique across the repo and this run.")
    st.dataframe(plan, width="stretch")
    if blocks:
        skipped = len(plan) - len(blocks)
        st.success(f"{len(blocks)} scan(s) ready." + (f" {skipped} skipped." if skipped else ""))
    else:
        st.warning("No scans passed validation — see the report above.")

    target_path = os.path.join(governance_dir, output_filename)
    file_exists, existing_text = read_target(target_path)
    show_target(environment, output_filename, target_path, file_exists, "dataplex-dq")
    render_preview_and_deploy(
        existing_text,
        render_file_header("dataplex-dq", scan_cron, catalog_publishing_enabled, export_dataset),
        blocks, target_path, governance_dir, output_filename, file_exists,
        scan_ids=ids, top_key="dataplex-dq")
